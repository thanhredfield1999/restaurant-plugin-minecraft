from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


OUTPUT = Path(__file__).with_name("Minecraft_Restaurant_Tycoon_Phase_4_Production_Hardening.xlsx")
STATUS = ["Chưa làm", "Đang thiết kế", "Sẵn sàng", "Đang làm", "Đang kiểm thử", "Bị chặn", "Hoàn thành", "Hủy"]
HEADERS = [
    "ID", "Nhóm công việc", "Tính năng/Hạng mục", "Mô tả cần điền", "Player Value", "Phụ thuộc",
    "Độ ưu tiên", "Độ phức tạp", "Người phụ trách", "Trạng thái", "Tiến độ %",
    "Acceptance Criteria", "Ghi chú/Evidence",
]


def task(group, name, description, value, dependency, priority, size, acceptance):
    return (group, name, description, value, dependency, priority, size, acceptance)


ROWS = [
    task("1. Launch Scope & Governance", "Production launch scope", "Chốt feature/content/platform/locale được launch, non-goal, feature flag và owner; điền launch date [ĐIỀN].", "Biết rõ trải nghiệm nào được hỗ trợ chính thức.", "Phase 3 exit report", "P0", "M", "Manifest scope được Product, Engineering và Operations ký; mọi mục có owner/trạng thái; artifact không chứa feature ngoài allowlist."),
    task("1. Launch Scope & Governance", "Supported platform matrix", "Chốt Minecraft/Paper/Java/client version, Việt/Anh, resource-pack policy và quyết định Geyser/Bedrock [ĐIỀN].", "Người chơi nhận thông tin tương thích rõ ràng.", "Production launch scope", "P0", "M", "100% tổ hợp được hỗ trợ pass login, pack, core loop, reconnect; ngoài matrix nhận thông báo unsupported rõ ràng."),
    task("1. Launch Scope & Governance", "Go/no-go authority", "Chỉ định release manager, approver Product/Engineering/QA/Operations và quyền dừng launch.", "Không phát hành build có rủi ro chưa được chấp thuận.", "Production launch scope", "P0", "S", "RACI được ký; tabletop một tình huống no-go cho đúng quyết định, escalation và audit trong SLA [ĐIỀN]."),
    task("1. Launch Scope & Governance", "Change freeze", "Định nghĩa mốc content/code/config freeze, exception process và emergency change.", "Giảm regression sát ngày mở cửa.", "Go/no-go authority", "P1", "S", "Mọi thay đổi sau freeze có ticket, risk, approver và regression evidence; unauthorized production drift = 0."),

    task("2. Compliance & Trust", "Minecraft EULA/Usage Guidelines review", "Review lại guideline tại thời điểm launch, disclaimer, operator identity/contact và cách quảng bá.", "Server minh bạch và giảm nguy cơ bị gián đoạn.", "Production launch scope", "P0", "L", "Checklist theo URL/phiên bản/ngày review có approver; 0 vi phạm mở; disclaimer và operator contact xuất hiện đúng nơi."),
    task("2. Compliance & Trust", "Branding and IP review", "Review tên, logo, screenshot, restaurant theme, UGC và quy trình IP takedown.", "Nội dung đáng tin cậy và có nguồn gốc hợp lệ.", "Asset provenance manifest", "P0", "M", "100% asset/brand claim có owner/license hoặc bị loại; takedown flow tabletop pass và có SLA [ĐIỀN]."),
    task("2. Compliance & Trust", "Asset provenance manifest", "Lập manifest source, creator, license, version/hash cho model, texture, sound, font, schematic và dependency.", "Tránh asset biến mất hoặc tranh chấp sau launch.", "Resource-pack release artifact", "P0", "L", "100% artifact production map tới provenance/license; unknown/incompatible license = 0; manifest hash gắn release."),
    task("2. Compliance & Trust", "Privacy and retention review", "Lập data inventory, purpose, lawful/consent basis nếu cần, access, TTL, export/delete và log redaction [ĐIỀN].", "Dữ liệu người chơi được thu tối thiểu và bảo vệ.", "Phase 3 analytics; moderation data", "P0", "L", "Mỗi field có purpose/owner/TTL; access, deletion và expiry tests pass; secret/PII fixture không xuất hiện trong log/analytics."),
    task("2. Compliance & Trust", "Moderation and child-safety readiness", "Chốt report, mute/block/ban, appeal, restaurant name/sign/chat review và escalation nghiêm trọng.", "Môi trường chơi an toàn và có hỗ trợ.", "Phase 3 moderation tools", "P0", "L", "Report-to-resolution tabletop pass; permission least-privilege; audit đầy đủ; on-call moderation roster đã xác nhận."),
    task("2. Compliance & Trust", "Monetization launch gate", "Quyết định [ĐIỀN: không launch/launch]; nếu launch, review price disclosure, history, refund/support, all-ages, no paid random và no competitive advantage.", "Mua hàng minh bạch và không pay-to-win.", "Policy review; payment design nếu có", "P0", "XL", "Nếu tắt: mọi store/payment path inaccessible. Nếu bật: legal/policy/security/refund/reconciliation tests pass và 0 đường mua tăng competitive power."),

    task("3. Release Engineering", "Immutable release artifact", "Build plugin/config/content/migration/resource pack thành artifact bất biến, có version, checksum và commit.", "Mọi server nhận đúng bản đã kiểm thử.", "CI pipeline; scope manifest", "P0", "L", "Exact artifact qua dev→staging→canary không rebuild; checksum khớp 100%; SBOM/provenance và signer [ĐIỀN] được lưu."),
    task("3. Release Engineering", "Dependency and SBOM review", "Pin phiên bản, lập SBOM, license/vulnerability review và policy update dependency.", "Giảm lỗi và rủi ro chuỗi cung ứng.", "Immutable release artifact", "P0", "L", "100% production dependency được pin và có license; 0 vulnerability vượt policy [ĐIỀN] không có waiver/owner/deadline."),
    task("3. Release Engineering", "Migration preflight", "Kiểm tra schema version, compatibility, backup restore point, duration/lock và forward-fix/rollback decision.", "Update không làm hỏng tiến trình.", "Release artifact; DB copy", "P0", "XL", "Fresh install và production-like upgrade pass; row count/checksum/invariant khớp; migration failure dừng deploy an toàn."),
    task("3. Release Engineering", "Deployment automation", "Tự động hóa precheck, maintenance/drain, migration, deploy, health, smoke, promote và abort.", "Giảm downtime và sai thao tác thủ công.", "Immutable artifact; migration preflight", "P0", "XL", "Staging dry-run hoàn tất không bước tay ngoài approval; lỗi inject tại mỗi stage dừng/rollback đúng và tạo audit."),
    task("3. Release Engineering", "Configuration and secret management", "Tách config/secret, validate schema, least privilege, rotation và chống log secret.", "Bảo vệ tài khoản và giữ cấu hình nhất quán.", "Deployment automation", "P0", "L", "Secret không nằm repo/artifact/log; invalid config fail-fast; rotation rehearsal không mất giao dịch và access cũ bị revoke."),
    task("3. Release Engineering", "Canary rollout", "Chọn scope [ĐIỀN], thời gian quan sát, metric gate và promote/rollback authority.", "Hạn chế người chơi bị ảnh hưởng bởi bản lỗi.", "Deployment; observability; rollback", "P0", "L", "Canary nhận đúng artifact/config/pack; synthetic gate pass; breach ngưỡng [ĐIỀN] tự dừng promote và rollback trong RTO."),
    task("3. Release Engineering", "Production rollback automation", "Rollback plugin/config/content/pack và chọn rollback/forward-fix DB; giữ correlation/audit.", "Sự cố release được đảo nhanh, có kiểm soát.", "Canary; backup; reconciliation", "P0", "XL", "Diễn tập N→N-1 trên production-like đạt RTO [ĐIỀN]; smoke, schema compatibility và reconciliation pass trước reopen."),
    task("3. Release Engineering", "Release smoke suite", "Tự động kiểm tra health, login, claim, order, purchase, save/reconnect, admin và pack trên exact artifact.", "Ngăn mở server khi core loop đã hỏng.", "Deployment automation", "P0", "L", "100% P0 smoke pass trên exact checksum; fail bất kỳ case nào chặn promote; report gắn build/config/schema/pack."),

    task("4. Observability & Incident Response", "Production health dashboard", "Dashboard MSPT/TPS, CPU, heap/GC, entity/chunk/task, DB pool/latency, error và concurrency.", "Lag và lỗi được phát hiện trước khi lan rộng.", "Production telemetry", "P0", "L", "Synthetic load/failure hiện đúng panel trong delay [ĐIỀN]; p50/p95/p99/max có build tag; critical No Data tạo cảnh báo."),
    task("4. Observability & Incident Response", "Economy and data-integrity dashboard", "Theo dõi earn/spend, balance-ledger delta, duplicate rejection, pending operation, reconcile mismatch và correction.", "Bảo vệ tiền và tiến trình người chơi.", "Ledger; reconciliation metrics", "P0", "L", "Fixture duplicate/mismatch tạo metric/alert; aggregate khớp ledger; unknown reason = 0; dashboard drill-down bằng correlation ID."),
    task("4. Observability & Incident Response", "SLO and alert policy", "Chốt SLI/SLO [ĐIỀN], burn-rate/threshold, severity, page owner và chống alert noise.", "Sự cố ảnh hưởng thật được xử lý đúng mức.", "Dashboards; capacity evidence", "P0", "L", "Mỗi alert test fire/resolution pass; page đến đúng owner trong SLA; dưới ngưỡng không page; runbook link hợp lệ."),
    task("4. Observability & Incident Response", "Structured logs and correlation", "Chuẩn hóa build, server, player pseudonymous ID, restaurant/run/operation/correlation ID và scrub policy.", "Support tìm nguyên nhân nhanh mà không lộ dữ liệu.", "Privacy review", "P0", "M", "Scripted transaction truy vết end-to-end; 100% critical event có correlation; secret/PII fixtures bị scrub."),
    task("4. Observability & Incident Response", "Incident severity and on-call", "Chốt SEV matrix, primary/secondary, escalation, break-glass và response target [ĐIỀN].", "Người chơi nhận hỗ trợ nhất quán khi có sự cố.", "SLO; moderation roster", "P0", "M", "Roster được xác nhận; paging drill primary-failover pass; break-glass được audit/revoke; scenario phân loại nhất quán."),
    task("4. Observability & Incident Response", "Incident runbooks", "Runbook database down, lag, duplication, stuck plot, pack/CDN, deploy, rollback và data loss.", "Giảm thời gian gián đoạn và thao tác nguy hiểm.", "Alerts; admin tools", "P0", "L", "Mỗi P0/P1 alert có runbook owner; tabletop từ detect→contain→recover→communicate pass, không sửa DB ad-hoc."),
    task("4. Observability & Incident Response", "Status communication", "Chuẩn bị maintenance, investigating, identified, monitoring, resolved và correction templates.", "Người chơi biết tác động và thời điểm cập nhật.", "Incident runbooks", "P1", "S", "Tabletop gửi đúng audience/cadence [ĐIỀN], nêu impact/action/update time và không hứa điều chưa xác nhận."),

    task("5. Backup, Restore & Reconciliation", "RPO/RTO contract", "Chốt RPO/RTO riêng cho DB, world, content/config và toàn dịch vụ [ĐIỀN].", "Mức mất dữ liệu/downtime được kiểm soát rõ.", "Business/product decision", "P0", "M", "Mỗi data class có RPO/RTO, owner và measurement method; Product/Engineering/Operations ký, không còn mục production TBD."),
    task("5. Backup, Restore & Reconciliation", "Database backup and PITR", "Cấu hình encrypted off-host backup và đánh giá WAL/PITR PostgreSQL; retention [ĐIỀN].", "Tiền và unlock có thể phục hồi đến điểm nhất quán.", "RPO/RTO; PostgreSQL", "P0", "XL", "Backup đạt lịch/RPO; restore tới timestamp mẫu pass checksum/invariant; job failure page đúng owner; access test pass."),
    task("5. Backup, Restore & Reconciliation", "World/content/config backup", "Version world snapshot, canonical stage, content/config/artifact và pack compatibility metadata.", "Nhà hàng phục hồi đúng hình ảnh tương ứng dữ liệu.", "Immutable artifacts; RPO/RTO", "P0", "L", "Mỗi backup map được DB point và version set; missing/mismatch bị preflight chặn; restore không dùng WorldEdit undo history."),
    task("5. Backup, Restore & Reconciliation", "Coordinated restore drill", "Restore vào môi trường cô lập, chọn DB/world/version point, chạy migration/reconcile/smoke rồi mới cutover.", "Công sức người chơi không mất hoặc nhân đôi sau thảm họa.", "All backups; runbooks", "P0", "XL", "Drill đạt RPO/RTO [ĐIỀN]; balance/unlock/order/prestige/world invariants sạch; 0 duplicate; biên bản và evidence được ký."),
    task("5. Backup, Restore & Reconciliation", "Post-restore reconciliation", "Đối chiếu ledger/balance, unlock/stage, operation fence, order/entitlement, prestige, entity và resource version.", "Phục hồi không để trạng thái ngầm bị sai.", "Restore drill; reconciliation tooling", "P0", "XL", "100% fixture sai được phát hiện; rerun idempotent; unresolved mismatch = 0 trước reopen hoặc có no-go decision."),
    task("5. Backup, Restore & Reconciliation", "Backup security and retention", "Kiểm tra encryption, least privilege, key rotation, deletion, immutability và restore access.", "Bản sao lưu không trở thành điểm rò rỉ dữ liệu.", "Privacy; backup systems", "P0", "L", "Unauthorized access bị deny/audit; key rotation và retention expiry pass; backup production không tải về máy cá nhân."),

    task("6. Resilience & Failure Injection", "Long soak test", "Chạy workload production-like [ĐIỀN: duration] gồm join/leave, plot churn, NPC, purchase, co-op và visit.", "Server ổn định qua nhiều giờ/ngày hoạt động.", "Capacity model; dashboards", "P0", "XL", "Không crash/deadlock/invariant violation; heap/entity/task/chunk/connection slope dưới ngưỡng [ĐIỀN]; restart sau soak pass."),
    task("6. Resilience & Failure Injection", "Reconnect storm", "Mô phỏng [ĐIỀN] client disconnect/reconnect đồng thời khi plot/order/purchase đang hoạt động.", "Mất mạng diện rộng không làm server hoặc dữ liệu hỏng.", "Load generator; session recovery", "P0", "L", "Login/session queues bounded; tick/DB đạt threshold; 0 duplicate session/reward/entitlement; recovery trong [ĐIỀN]."),
    task("6. Resilience & Failure Injection", "Degraded database test", "Inject latency, timeout, disconnect, pool exhaustion, read-only và recovery.", "Database lỗi không làm sai economy.", "Degraded-mode policy; observability", "P0", "XL", "Economic write fail-closed; main thread không block; committed operation không mất/double; exit degraded mode reconciliation sạch."),
    task("6. Resilience & Failure Injection", "Crash injection matrix", "Kill process tại mỗi phase purchase, payout, reset, prestige, inventory entitlement và world operation.", "Crash không làm mất tiền, nhân đôi món hoặc kẹt plot.", "Durable state machines; recovery", "P0", "XL", "Mỗi checkpoint chạy [ĐIỀN] lần; operation đạt terminal/retryable state; ledger delta 0; canonical world và entitlement khớp."),
    task("6. Resilience & Failure Injection", "Resource-pack/CDN failure", "Inject reject, timeout, invalid URL/hash, corrupt cache, failed reload, CDN outage và rollback version.", "Người chơi không gặp control vô hình hoặc vòng tải vô hạn.", "Supported client matrix; pack artifact", "P0", "L", "Mọi client support theo đúng required/optional policy; retry/help/fallback rõ; pack mismatch không vào gameplay phụ thuộc asset."),
    task("6. Resilience & Failure Injection", "Capacity and overload test", "Chạy expected/peak/beyond-peak theo active plot, NPC, path, construction, DB và pack workload [ĐIỀN].", "Trải nghiệm mượt trong tải công bố và giảm tải an toàn khi quá mức.", "Hardware; budgets; SLO", "P0", "XL", "Expected/peak đạt MSPT p95 dưới budget và <50 ms; overload kích hoạt backpressure, không crash/corrupt/unbounded growth."),
    task("6. Resilience & Failure Injection", "Plot reset/reassign chaos", "Trộn reset, reassign, unload/reload, visitor evacuation, purchase và stale callback.", "Không có NPC ma, sửa nhầm plot hoặc mất progression.", "Fence token; entity lifecycle", "P0", "XL", "Stale epoch/fence mutation = 0; orphan entity/task/chunk lease = 0; canonical stage/revision hội tụ sau mọi vòng."),

    task("7. Launch Readiness", "Production security review", "Review permissions, admin commands, DB/network, secrets, plugin surface, rate limit và audit tamper/access.", "Giảm lạm dụng quyền và phá hoại server.", "Compliance; config; admin tools", "P0", "L", "Negative permission tests pass; least privilege; security findings vượt policy = 0 hoặc no-go waiver được ký."),
    task("7. Launch Readiness", "Operational access rehearsal", "Xác nhận dashboard/log/DB/host/CDN/backup/status access cho đúng role, MFA và break-glass.", "On-call có thể xử lý sự cố mà không chia sẻ tài khoản.", "On-call roster; security review", "P0", "M", "Primary/secondary thực hiện read-only diagnosis và approved recovery; shared credential = 0; revoke test pass."),
    task("7. Launch Readiness", "Support readiness", "Chuẩn bị FAQ, known issues, report/support route, refund/correction policy và triage queue.", "Người chơi biết cách nhận trợ giúp.", "Scope; moderation; communications", "P1", "M", "Sample ticket từ submit→triage→resolution pass; SLA/owner rõ; không yêu cầu secret hoặc dữ liệu thừa."),
    task("7. Launch Readiness", "Launch-day runbook", "Timeline T-24h đến T+24h, freeze, backup, deploy, canary, promote, monitoring, comms và rollback triggers.", "Ngày mở server có trách nhiệm và quyết định rõ.", "All production runbooks", "P0", "L", "Tabletop đầy đủ vai trò pass; mọi bước có owner/evidence; trigger no-go/rollback và communication cadence [ĐIỀN] rõ."),
    task("7. Launch Readiness", "Final production acceptance", "Chạy exact release candidate qua compliance, security, smoke, regression, exploit, failure, restore và operational review.", "Chỉ phát hành phiên bản có bằng chứng chất lượng đầy đủ.", "All P0 tasks", "P0", "XL", "100% release blocker hoàn thành; 0 P0/P1 không được chấp thuận; exact hashes được ký bởi Product, Engineering, QA và Operations."),
    task("7. Launch Readiness", "Post-launch review", "Lên lịch checkpoint [ĐIỀN], review SLO, incident, economy, progression, moderation và capacity; không mở scope trước review.", "Sự cố thực tế được chuyển thành cải tiến có kiểm soát.", "Production launch", "P1", "S", "Review có metric/evidence, action owner/deadline và quyết định giữ/rollback/tune; không có finding nghiêm trọng vô chủ."),
]


MILESTONES = [
    ("M1", "Scope & trust lock", "[ĐIỀN: ngày]", "Launch scope/platform, compliance, privacy, moderation, asset/dependency provenance được ký.", "Chưa làm"),
    ("M2", "Release pipeline ready", "[ĐIỀN: ngày]", "Artifact, migration, deploy, canary, rollback, smoke và secret management drill pass.", "Chưa làm"),
    ("M3", "Recovery proven", "[ĐIỀN: ngày]", "RPO/RTO, backup/PITR, coordinated restore và reconciliation drill pass.", "Chưa làm"),
    ("M4", "Resilience proven", "[ĐIỀN: ngày]", "Soak, reconnect, DB degraded, crash, pack/CDN, overload và plot chaos pass.", "Chưa làm"),
    ("M5", "Production launch gate", "[ĐIỀN: ngày]", "Final acceptance ký; 0 release blocker; launch-day runbook tabletop pass.", "Chưa làm"),
]

RISKS = [
    ("R-01", "Restore DB/world lệch thời điểm tạo duplicate hoặc mất stage", "Rất cao", "Coordinated point, immutable versions, restore cô lập và reconciliation bắt buộc."),
    ("R-02", "Canary không đại diện tải production", "Cao", "Synthetic workload + cohort thật, metric gate và staged ramp theo active plot/NPC."),
    ("R-03", "Crash/failure test bỏ sót transaction boundary", "Rất cao", "State-transition inventory, deterministic fault injection và invariant scan mỗi checkpoint."),
    ("R-04", "License/policy finding sát launch", "Cao", "Freeze asset/dependency sớm; manifest và approver; có phương án loại asset/store."),
    ("R-05", "Alert noise hoặc thiếu signal làm chậm incident", "Cao", "Test fire/resolution, SLO-based paging, dashboard No Data và tabletop on-call."),
    ("R-06", "Content scope creep thay đổi release candidate", "Cao", "Change freeze, feature flag, exact artifact hash; cắt content trước hardening."),
    ("R-07", "Solo operator không đáp ứng on-call/moderation", "Cao", "Giảm launch window/concurrency, backup contact và công khai support hours."),
]

EXIT_GATES = [
    ("PG-01", "Scope & platform", "Scope/version/locale/pack/Geyser decision được ký; support matrix pass."),
    ("PG-02", "Compliance & trust", "EULA/usage, branding/IP, asset/dependency license, privacy, moderation và monetization gate pass."),
    ("PG-03", "Release safety", "Immutable artifact, migration, deploy/canary/rollback và exact-artifact smoke pass."),
    ("PG-04", "Data durability", "Backup/PITR, coordinated restore đạt RPO/RTO và reconciliation không còn mismatch."),
    ("PG-05", "Observability", "Dashboard, SLO/alerts, correlation, on-call và runbook drills pass."),
    ("PG-06", "Resilience", "Long soak, reconnect storm, degraded DB, crash injection, pack/CDN và overload pass."),
    ("PG-07", "Security & operations", "Permission/secret/access review và launch-day tabletop pass."),
    ("PG-08", "Final sign-off", "0 P0/P1 không được duyệt; Product, Engineering, QA và Operations ký exact hashes."),
]


def fill(cell, color):
    cell.fill = PatternFill("solid", fgColor=color)


def style_table(ws, start_row, title, headers, rows, colors):
    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=13)
    c = ws.cell(start_row, 1, title)
    c.font = Font(name="Arial", size=12, bold=True, color="FFFFFF")
    fill(c, colors["navy"])
    for col, header in enumerate(headers, 1):
        c = ws.cell(start_row + 1, col, header)
        c.font = Font(name="Arial", size=9, bold=True, color="FFFFFF")
        fill(c, colors["blue"])
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r, values in enumerate(rows, start_row + 2):
        for col, value in enumerate(values, 1):
            c = ws.cell(r, col, value)
            c.font = Font(name="Arial", size=9)
            fill(c, "FFFFFF" if r % 2 else colors["pale"])
            c.alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[r].height = 42
    return start_row + len(rows) + 3


def build():
    wb = Workbook()
    ws = wb.active
    ws.title = "Phase 4 - Production"
    colors = {"navy": "172A3A", "blue": "1F5A7A", "pale": "EAF3F8", "red": "C00000", "pale_red": "FCE4D6", "yellow": "FFF2CC", "green": "E2F0D9"}
    thin = Side(style="thin", color="B7C9D6")
    border = Border(top=thin, bottom=thin, left=thin, right=thin)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A7"
    ws.merge_cells("A1:M1")
    ws["A1"] = "MINECRAFT RESTAURANT TYCOON | PHASE 4 - PRODUCTION HARDENING"
    ws["A1"].font = Font(name="Arial", size=16, bold=True, color="FFFFFF")
    fill(ws["A1"], colors["navy"])
    ws.merge_cells("A2:M2")
    ws["A2"] = "Mục tiêu: chứng minh khả năng phát hành, phục hồi và vận hành production. Chỉ sửa ô [ĐIỀN], owner, trạng thái, tiến độ và evidence."
    ws["A2"].font = Font(name="Arial", size=10, italic=True, color="334E68")
    fill(ws["A2"], colors["pale"])
    legends = [
        (3, "Hướng dẫn", "Ô vàng là quyết định cần điền. Dòng đỏ là release blocker P0; không launch khi chưa Hoàn thành."),
        (4, "Ví dụ", "[ĐIỀN: RPO 15 phút] | Owner: Platform | Trạng thái: Đang kiểm thử | Tiến độ: 75 | Evidence: link drill"),
        (5, "Nguồn", "RESTAURANT_TYCOON_PLAN.md §17 Phase 4, §18 Acceptance Gates và §23 Binding Production Requirements."),
    ]
    for r, label, text in legends:
        ws.cell(r, 1, label).font = Font(name="Arial", size=9, bold=True, color="FFFFFF")
        fill(ws.cell(r, 1), colors["blue"])
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=13)
        ws.cell(r, 2, text).font = Font(name="Arial", size=9)
        fill(ws.cell(r, 2), colors["yellow"] if r == 4 else colors["pale"])
    for col, header in enumerate(HEADERS, 1):
        c = ws.cell(6, col, header)
        c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        fill(c, colors["blue"])
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border
    widths = [15, 28, 31, 54, 38, 35, 12, 13, 22, 18, 12, 72, 34]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width

    group_ids = {}
    for row_number, item in enumerate(ROWS, 7):
        group, name, description, value, dependency, priority, size, acceptance = item
        group_ids[group] = group_ids.get(group, 0) + 1
        group_number = int(group.split(".")[0])
        item_id = f"P4-{group_number:02d}-{group_ids[group]:02d}"
        blocker = priority == "P0"
        note = "RELEASE BLOCKER | [ĐIỀN: link evidence sau khi pass]" if blocker else "[ĐIỀN: quyết định/link evidence]"
        values = [item_id, group, name, description, value, dependency, priority, size, "[ĐIỀN: tên/role]", "Chưa làm", 0, acceptance, note]
        for col, value_cell in enumerate(values, 1):
            c = ws.cell(row_number, col, value_cell)
            c.font = Font(name="Arial", size=9, color=colors["red"] if blocker and col in (1, 7, 13) else "1F1F1F", bold=blocker and col in (1, 7, 13))
            cell_color = colors["pale_red"] if blocker else ("FFFFFF" if group_number % 2 else colors["pale"])
            if "[ĐIỀN" in str(value_cell) and not blocker:
                cell_color = colors["yellow"]
            fill(c, cell_color)
            c.alignment = Alignment(horizontal="center" if col in (1, 7, 8, 10, 11) else "left", vertical="top", wrap_text=True)
            c.border = border
        ws.row_dimensions[row_number].height = 76
    end = 6 + len(ROWS)
    ws.auto_filter.ref = f"A6:M{end}"

    dvs = [
        (DataValidation(type="list", formula1='"' + ",".join(STATUS) + '"', allow_blank=False), f"J7:J{end}"),
        (DataValidation(type="list", formula1='"P0,P1,P2,P3"', allow_blank=False), f"G7:G{end}"),
        (DataValidation(type="list", formula1='"XS,S,M,L,XL"', allow_blank=False), f"H7:H{end}"),
        (DataValidation(type="whole", operator="between", formula1="0", formula2="100", allow_blank=False), f"K7:K{end}"),
    ]
    for dv, target in dvs:
        ws.add_data_validation(dv)
        dv.add(target)
    ws.conditional_formatting.add(f"A7:M{end}", FormulaRule(formula=['$J7="Hoàn thành"'], fill=PatternFill("solid", fgColor=colors["green"])))
    ws.conditional_formatting.add(f"A7:M{end}", FormulaRule(formula=['$J7="Bị chặn"'], fill=PatternFill("solid", fgColor=colors["pale_red"])))

    dashboard = wb.create_sheet("Dashboard")
    dashboard.sheet_view.showGridLines = False
    dashboard.merge_cells("A1:F1")
    dashboard["A1"] = "PHASE 4 | PRODUCTION READINESS DASHBOARD"
    dashboard["A1"].font = Font(name="Arial", size=16, bold=True, color="FFFFFF")
    fill(dashboard["A1"], colors["navy"])
    metrics = [
        ("Tổng hạng mục", f"=COUNTA('Phase 4 - Production'!$A$7:$A${end})"),
        ("Đã hoàn thành", f'=COUNTIF(\'Phase 4 - Production\'!$J$7:$J${end},"Hoàn thành")'),
        ("Release blocker P0", f'=COUNTIF(\'Phase 4 - Production\'!$G$7:$G${end},"P0")'),
        ("P0 đã hoàn thành", f'=COUNTIFS(\'Phase 4 - Production\'!$G$7:$G${end},"P0",\'Phase 4 - Production\'!$J$7:$J${end},"Hoàn thành")'),
        ("P0 còn mở", "=B5-B6"),
        ("Tiến độ trung bình", f"=AVERAGE('Phase 4 - Production'!$K$7:$K${end})/100"),
        ("Go/No-Go", '=IF(B7=0,"GO - chờ sign-off","NO-GO")'),
    ]
    dashboard["A2"] = "Chỉ số"
    dashboard["B2"] = "Giá trị"
    for c in dashboard[2]:
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        fill(c, colors["blue"])
    for r, (label, formula) in enumerate(metrics, 3):
        dashboard.cell(r, 1, label).font = Font(name="Arial", size=10)
        dashboard.cell(r, 2, formula).font = Font(name="Arial", size=10, bold=True)
        fill(dashboard.cell(r, 1), colors["pale"] if r % 2 else "FFFFFF")
        fill(dashboard.cell(r, 2), colors["pale"] if r % 2 else "FFFFFF")
    dashboard["B8"].number_format = "0.0%"
    dashboard.column_dimensions["A"].width = 30
    dashboard.column_dimensions["B"].width = 24
    dashboard["A12"] = "Công thức dashboard lấy trực tiếp từ backlog. GO chỉ có nghĩa P0 đã hoàn thành; final sign-off và các quyết định [ĐIỀN] vẫn bắt buộc."
    dashboard.merge_cells("A12:F13")
    dashboard["A12"].alignment = Alignment(wrap_text=True, vertical="top")
    dashboard["A12"].font = Font(name="Arial", size=9, italic=True)
    fill(dashboard["A12"], colors["yellow"])

    row = end + 3
    row = style_table(ws, row, "Milestones", ["ID", "Milestone", "Ngày mục tiêu", "Điều kiện hoàn thành", "Trạng thái"], MILESTONES, colors)
    risk_rows = [(rid, risk, impact, mitigation, "[ĐIỀN: owner]", "Mở") for rid, risk, impact, mitigation in RISKS]
    row = style_table(ws, row, "Production Risks", ["ID", "Risk", "Ảnh hưởng", "Mitigation/Contingency", "Owner", "Trạng thái"], risk_rows, colors)
    gate_rows = [(gid, group, criteria, "[ĐIỀN: approver]", "Chưa làm") for gid, group, criteria in EXIT_GATES]
    row = style_table(ws, row, "Production Exit Gates", ["ID", "Nhóm tiêu chí", "Điều kiện Go", "Approver", "Kết quả"], gate_rows, colors)
    dod = [
        ("DoD-01", "Evidence", "Test/drill evidence gắn exact build/config/schema/content/pack hashes; có timestamp và operator."),
        ("DoD-02", "Failure behavior", "Happy, boundary, retry/concurrent, restart/crash, degraded và rollback path liên quan đều pass."),
        ("DoD-03", "Data", "Không duplicate/loss; ledger và durable state khớp; reconciliation rerun idempotent và sạch."),
        ("DoD-04", "Operations", "Dashboard/alert/runbook/owner hoạt động; restore/rollback path được drill nếu hạng mục có impact."),
        ("DoD-05", "Trust", "Privacy, least privilege, audit, license/provenance và player communication được review khi áp dụng."),
        ("DoD-06", "Sign-off", "0 P0/P1 mở không waiver; approver ký và mọi [ĐIỀN] bắt buộc có quyết định."),
    ]
    style_table(ws, row, "Definition of Done", ["ID", "Khía cạnh", "Điều kiện Done"], dod, colors)
    ws.print_title_rows = "1:6"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.save(OUTPUT)
    return end


def verify(end):
    wb = load_workbook(OUTPUT, data_only=False)
    ws = wb["Phase 4 - Production"]
    assert end - 6 == len(ROWS)
    assert len({ws.cell(r, 1).value for r in range(7, end + 1)}) == len(ROWS)
    assert all(ws.cell(r, 12).value and len(ws.cell(r, 12).value) > 50 for r in range(7, end + 1))
    assert all(ws.cell(r, 13).value.startswith("RELEASE BLOCKER") for r in range(7, end + 1) if ws.cell(r, 7).value == "P0")
    assert "Dashboard" in wb.sheetnames
    dashboard = wb["Dashboard"]
    assert dashboard["B3"].data_type == "f"
    assert dashboard["B7"].value == "=B5-B6"
    assert dashboard["B9"].value == '=IF(B7=0,"GO - chờ sign-off","NO-GO")'
    assert len(dashboard.formula_attributes) == 0 if hasattr(dashboard, "formula_attributes") else True
    assert len(ws.data_validations.dataValidation) == 4
    print(f"Created: {OUTPUT}")
    print(f"Plan items: {len(ROWS)}; P0 blockers: {sum(row[5] == 'P0' for row in ROWS)}; sheets: {len(wb.sheetnames)}")


if __name__ == "__main__":
    verify(build())
