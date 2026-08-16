from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


OUTPUT = Path(__file__).with_name("Minecraft_Restaurant_Tycoon_Phase_3_Closed_Beta.xlsx")

GROUPS = [
    ("Beta Scope", [
        "Mục tiêu Closed Beta", "Số lượng người chơi thử nghiệm", "Thời gian beta",
        "Tiêu chí chọn tester", "Nội dung được phép phát hành", "Nội dung bị loại khỏi beta",
        "Exit criteria",
    ]),
    ("Content Expansion", [
        "Theme thứ hai hoặc prestige route", "Build stage mới", "Purchase node mới", "Recipe mới",
        "Customer type mới", "Staff/automation upgrade", "Cosmetic reward", "Visual và sound polish",
    ]),
    ("Economy Tuning", [
        "Early-game pacing", "Mid-game pacing", "First prestige pacing", "Purchase price", "Revenue curve",
        "Manual/automation balance", "Currency source và sink", "Inflation monitoring", "Economy rollback plan",
    ]),
    ("Progression Tuning", [
        "Progression bottleneck", "Hai progression branch", "Unlock timing", "Tutorial completion",
        "Session pacing", "Retention objective", "Prestige conversion", "Catch-up rule nếu có",
    ]),
    ("Seasonal Objectives", [
        "Season theme", "Objective", "Điều kiện tham gia", "Progress tracking", "Reward", "Anti-AFK",
        "Anti-pay-to-win", "Reset khi hết season", "Recovery khi restart", "Seasonal analytics",
    ]),
    ("Exploit Testing", [
        "Duplicate item", "Duplicate payout", "Inventory click/drag", "Shift-click", "Number-key swap",
        "Double-click", "Disconnect trong transaction", "Restart trong transaction",
        "Co-op reward duplication", "Visitor permission bypass", "Purchase/reset race condition",
        "Prestige duplication",
    ]),
    ("Reliability", [
        "Server restart", "Plugin reload policy", "Database outage", "Database latency", "Crash recovery",
        "Plot reconciliation", "Entity cleanup", "Stuck customer recovery", "Pending transaction recovery",
        "Degraded mode",
    ]),
    ("Rollback", [
        "Plugin rollback", "Database migration rollback", "Content rollback", "Resource-pack rollback",
        "Economy correction", "Player-data reconciliation", "Rollback verification",
        "Communication khi rollback",
    ]),
    ("Resource Pack", [
        "Download flow", "Versioning", "Hash validation", "Người chơi từ chối pack", "Download failure",
        "CDN failure", "Fallback presentation", "Client compatibility", "Rollback pack version",
    ]),
    ("Multiplayer và Moderation", [
        "Co-op permission", "Visitor permission", "Grief prevention", "Chat moderation", "Report flow",
        "Player sanction", "Audit log", "Admin escalation", "Privacy rule",
    ]),
    ("Operations", [
        "Deployment checklist", "Canary release", "Monitoring dashboard", "Alert rule", "Incident severity",
        "Incident response", "Backup schedule", "Restore procedure", "Maintenance mode",
        "Player communication",
    ]),
    ("Analytics và Feedback", [
        "Error tracking", "Performance tracking", "Economy tracking", "Progression tracking",
        "Retention tracking", "Tester feedback form", "Bug report template", "Feedback prioritization",
        "Privacy và data retention",
    ]),
    ("Performance", [
        "Expected beta concurrency", "Peak concurrency test", "Active plot count", "NPC count",
        "Entity tracking", "Pathfinding load", "Tick time", "Memory", "Database pool", "Long soak test",
    ]),
    ("Beta Testing", [
        "Smoke test", "Regression test", "Gameplay test", "Economy test", "Exploit test", "Restart test",
        "Rollback test", "Resource-pack failure test", "Backup/restore test", "Moderation test",
        "Release acceptance test",
    ]),
]

DESCRIPTIONS = {
    "Mục tiêu Closed Beta": "Chốt mục tiêu học hỏi, phạm vi vận hành thật và các giả thuyết cần xác nhận trong Closed Beta.",
    "Số lượng người chơi thử nghiệm": "Điền quy mô tester mục tiêu, concurrency dự kiến và cách chia cohort; không coi đây là capacity đã được chứng minh.",
    "Thời gian beta": "Điền ngày bắt đầu, ngày kết thúc, khung bảo trì và các checkpoint đánh giá.",
    "Tiêu chí chọn tester": "Định nghĩa tiêu chí tuyển chọn, mức độ đa dạng thiết bị/khu vực/playstyle và NDA nếu áp dụng.",
    "Nội dung được phép phát hành": "Liệt kê feature, content version và config được phép xuất hiện trong build beta.",
    "Nội dung bị loại khỏi beta": "Liệt kê feature flag bị tắt, nội dung hoãn và hành vi không được quảng bá.",
    "Exit criteria": "Chốt ngưỡng chất lượng, ổn định, economy, performance và vận hành để đóng beta hoặc chuyển phase.",
    "Theme thứ hai hoặc prestige route": "Chọn đúng một phương án mở rộng: [ĐIỀN: theme thứ hai] hoặc [ĐIỀN: prestige route]; mô tả fantasy và phạm vi tối thiểu.",
    "Build stage mới": "Thiết kế [ĐIỀN: build stage], thay đổi world projection và giá trị gameplay tương ứng.",
    "Purchase node mới": "Định nghĩa [ĐIỀN: purchase node], prerequisite, hiệu ứng, idempotency key và presentation.",
    "Recipe mới": "Định nghĩa [ĐIỀN: recipe], station sequence, thời gian xử lý và recipe/version ID.",
    "Customer type mới": "Định nghĩa [ĐIỀN: customer type], hành vi khác biệt, timeout và economic result có giới hạn.",
    "Staff/automation upgrade": "Định nghĩa [ĐIỀN: staff/automation upgrade], công việc được tự động hóa và giới hạn throughput.",
    "Cosmetic reward": "Định nghĩa [ĐIỀN: cosmetic reward] chỉ thay đổi presentation, không cấp lợi thế economy/gameplay.",
    "Visual và sound polish": "Polish feedback mua hàng, unlock, phục vụ, lỗi và trạng thái degraded bằng visual/sound nhất quán.",
    "Early-game pacing": "Dùng analytics Alpha để đề xuất target pacing early-game và các config cần tune; không hardcode quyết định trước sign-off.",
    "Mid-game pacing": "Xác định funnel, thời gian chờ và khoảng cách purchase trong mid-game từ dữ liệu cohort Alpha.",
    "First prestige pacing": "Xác định target time-to-first-prestige theo cohort và điều kiện hợp lệ của session.",
    "Purchase price": "Rà soát giá từng purchase node theo dependency graph, earning rate và target pacing được duyệt.",
    "Revenue curve": "Mô hình hóa revenue theo stage, active plot và automation; phát hiện spike, plateau hoặc runaway growth.",
    "Manual/automation balance": "Tune phần thưởng manual so với automation theo cửa sổ đo [ĐIỀN: window] và trần uplift được duyệt.",
    "Currency source và sink": "Lập inventory mọi source/sink, ownership, idempotency và dashboard đối soát.",
    "Inflation monitoring": "Theo dõi balance distribution, earning/spend rate và cảnh báo theo ngưỡng [ĐIỀN: thresholds].",
    "Economy rollback plan": "Định nghĩa versioned economy config, phạm vi revert, correction ledger và phê duyệt rollback.",
    "Progression bottleneck": "Tìm bottleneck từ funnel, time-to-unlock và feedback; phân biệt friction chủ ý với lỗi chặn.",
    "Hai progression branch": "Xác nhận hai branch đều hợp lệ, có trade-off dễ hiểu và không tạo dead end.",
    "Unlock timing": "Tune thời điểm unlock theo event analytics và target pacing [ĐIỀN: target].",
    "Tutorial completion": "Đo completion/drop-off từng bước tutorial và sửa bước gây thất bại không chủ ý.",
    "Session pacing": "Định nghĩa nhịp mục tiêu cho session [ĐIỀN: duration bands] và khoảng cách meaningful event.",
    "Retention objective": "Điền mục tiêu retention, cohort definition và cửa sổ đo; không dùng reward pay-to-win để ép quay lại.",
    "Prestige conversion": "Đo tỷ lệ player đủ điều kiện thực hiện prestige và nguyên nhân từ chối/churn.",
    "Catch-up rule nếu có": "Quyết định [ĐIỀN: có/không]; nếu có, giới hạn catch-up để không phá economy hoặc lợi thế cạnh tranh.",
    "Season theme": "Điền concept, thời gian và content boundary cho [ĐIỀN: season theme].",
    "Objective": "Định nghĩa [ĐIỀN: objective] dựa trên gameplay hợp lệ, có progress unit và completion rule rõ ràng.",
    "Điều kiện tham gia": "Điền eligibility, opt-in, prerequisite và xử lý player tham gia giữa season.",
    "Progress tracking": "Thiết kế server-authoritative progress event, idempotency và hiển thị tiến độ.",
    "Reward": "Điền [ĐIỀN: reward]; kiểm tra reward không tăng earning power, unlock power hoặc tạo lợi thế trả phí.",
    "Anti-AFK": "Loại trừ idle/repeated synthetic action khỏi seasonal progress bằng rule có thể audit.",
    "Anti-pay-to-win": "Kiểm kê mọi đường mua/nhận lợi thế và chứng minh thanh toán không rút ngắn objective hoặc tăng power.",
    "Reset khi hết season": "Định nghĩa close/freeze/archive/reset state và xử lý reward chưa claim.",
    "Recovery khi restart": "Bảo toàn seasonal progress qua graceful restart và crash, không mất hoặc cộng đôi.",
    "Seasonal analytics": "Ghi event participation, progress, completion, abandonment và reward claim theo version.",
    "Plugin reload policy": "Cấm hoặc giới hạn plugin reload; quy định restart được hỗ trợ và thông điệp cho operator.",
    "Server restart": "Kiểm tra graceful restart tại mọi state quan trọng của order, purchase, build và season.",
    "Database outage": "Thiết kế hành vi khi database unavailable: chặn write nguy hiểm, bảo toàn state và phục hồi có kiểm soát.",
    "Database latency": "Định nghĩa timeout, backpressure và UX khi DB chậm mà không block main thread.",
    "Crash recovery": "Khôi phục state committed từ database và reconcile world projection sau process crash.",
    "Plot reconciliation": "So sánh durable state với block/entity projection và repair idempotently.",
    "Entity cleanup": "Dọn entity/task còn sót khi plot inactive, unload, reset hoặc plugin shutdown.",
    "Stuck customer recovery": "Watchdog customer state/path timeout, giải phóng reservation và không phạt player.",
    "Pending transaction recovery": "Quét transaction pending/unknown, quyết định resume/compensate theo idempotency key.",
    "Degraded mode": "Định nghĩa feature được giữ/tắt khi dependency lỗi và điều kiện tự động thoát degraded mode.",
    "Plugin rollback": "Định nghĩa artifact pin, config compatibility và quy trình quay lại plugin build trước.",
    "Database migration rollback": "Mỗi migration beta có precheck, rollback/forward-fix được diễn tập và giới hạn dữ liệu rõ ràng.",
    "Content rollback": "Version hóa content/config để tắt hoặc revert node/recipe/theme không làm hỏng save.",
    "Resource-pack rollback": "Phối hợp server config và pack version tương thích khi rollback.",
    "Economy correction": "Dùng correction ledger có lý do, scope, dry-run và audit; không sửa balance âm thầm.",
    "Player-data reconciliation": "Đối chiếu balance, unlock, prestige, entitlement và world projection sau rollback.",
    "Rollback verification": "Checklist xác minh schema, build, config, gameplay, data và monitoring sau rollback.",
    "Communication khi rollback": "Chuẩn bị template thông báo downtime, impact, correction và thời điểm cập nhật tiếp theo.",
    "Download flow": "Kiểm tra prompt, URL, consent, trạng thái tải và activation của resource pack.",
    "Versioning": "Gắn pack version bất biến với build/content compatibility matrix.",
    "Hash validation": "Cấu hình hash đúng artifact và từ chối cache/artifact không khớp.",
    "Người chơi từ chối pack": "Định nghĩa policy [ĐIỀN: allow/kick/degraded] và thông báo rõ trước khi vào gameplay.",
    "Download failure": "Cung cấp retry/fallback/support code khi client không tải được pack.",
    "CDN failure": "Diễn tập origin/CDN unavailable và kích hoạt fallback hoặc maintenance theo runbook.",
    "Fallback presentation": "Đảm bảo item/UI cốt lõi vẫn hiểu được khi custom model/sound không hoạt động.",
    "Client compatibility": "Xác nhận Minecraft client versions, locale và platform được hỗ trợ [ĐIỀN: matrix].",
    "Rollback pack version": "Cho phép chuyển về pack version trước mà client không bị vòng lặp tải hoặc mismatch.",
    "Co-op permission": "Lập permission matrix cho owner/member và kiểm tra mọi hành động kinh tế/world-side ở server.",
    "Visitor permission": "Visitor chỉ xem/tương tác được action allowlist; không mua, reset, claim hoặc nhận reward chủ plot.",
    "Grief prevention": "Chặn break/place, container, entity, fluid, piston và interaction ngoài quyền trong plot.",
    "Chat moderation": "Định nghĩa filter, mute/escalation, appeal và xử lý false positive theo policy.",
    "Report flow": "Cho phép gửi report có category, target, evidence metadata và confirmation ID.",
    "Player sanction": "Định nghĩa warning/mute/kick/ban, duration, reason code, appeal và dual-control nếu cần.",
    "Audit log": "Ghi actor, action, target, before/after, reason, timestamp và correlation ID cho thao tác nhạy cảm.",
    "Admin escalation": "Định nghĩa on-call/escalation path, quyền break-glass và thời gian phản hồi [ĐIỀN: SLA].",
    "Privacy rule": "Định nghĩa dữ liệu được xem trong moderation, least privilege và cách xử lý yêu cầu privacy.",
    "Deployment checklist": "Checklist artifact, migration, backup, config, pack, canary, smoke test và rollback owner.",
    "Canary release": "Phát hành cho cohort/server [ĐIỀN: canary scope], quan sát rồi promote/rollback theo gate.",
    "Monitoring dashboard": "Dashboard health, tick, memory, DB, error, transaction, plot và player-impact metrics.",
    "Alert rule": "Định nghĩa signal, threshold [ĐIỀN], duration, severity, owner và chống alert noise.",
    "Incident severity": "Chốt SEV matrix theo impact player/data/economy và response target [ĐIỀN].",
    "Incident response": "Runbook detect, triage, contain, recover, communicate và post-incident review.",
    "Backup schedule": "Điền RPO, lịch backup, retention, encryption, location và owner kiểm tra.",
    "Restore procedure": "Runbook restore vào môi trường cô lập, validate rồi mới quyết định cutover.",
    "Maintenance mode": "Chặn login/write có kiểm soát, giữ admin access và hiển thị ETA/message.",
    "Player communication": "Template status cho planned maintenance, incident, rollback và resolved update.",
    "Error tracking": "Capture exception có build/version/correlation context, scrub dữ liệu nhạy cảm và nhóm lỗi.",
    "Performance tracking": "Thu tick time, task duration, entity/pathfinding, DB latency/pool và memory theo build.",
    "Economy tracking": "Event ledger/aggregate cho earn, spend, balance, purchase, refund và prestige theo config version.",
    "Progression tracking": "Event claim, tutorial step, unlock, branch, build stage và prestige funnel.",
    "Retention tracking": "Định nghĩa cohort/timezone/return event và loại test/admin account khỏi metric.",
    "Tester feedback form": "Form thu build, context, severity cảm nhận, reproduction và consent liên hệ.",
    "Bug report template": "Template có environment, steps, expected/actual, evidence và impact/data-risk flag.",
    "Feedback prioritization": "Triage theo severity, frequency, player value, data risk và release blocker status.",
    "Privacy và data retention": "Liệt kê event/field, mục đích, retention [ĐIỀN], access và deletion process.",
    "Expected beta concurrency": "Điền expected/peak concurrency giả định và workload profile; chỉ công bố sau benchmark.",
    "Peak concurrency test": "Load test workload đại diện ở [ĐIỀN: peak] với ramp, steady-state và pass thresholds.",
    "Active plot count": "Đo capacity theo active/observed plot thay vì chỉ player count.",
    "NPC count": "Đặt budget [ĐIỀN] theo observed plot và xác nhận cleanup/inactive simulation.",
    "Entity tracking": "Đo tổng tracked entity, display/interaction entity và client-visible count theo plot.",
    "Pathfinding load": "Đo path request rate/cost; chỉ repath khi đổi state hoặc watchdog yêu cầu.",
    "Tick time": "Chốt ngưỡng p50/p95/p99 tick time hoặc MSPT [ĐIỀN] và đo dưới workload beta.",
    "Memory": "Đo heap/native memory, GC pause và growth; đặt ngưỡng [ĐIỀN] và kiểm tra leak.",
    "Database pool": "Tune pool/timeout theo concurrency [ĐIỀN], không block server main thread.",
    "Long soak test": "Chạy soak [ĐIỀN: duration] với traffic, restart/failure injection và leak checks.",
}

TEST_DESCRIPTIONS = {
    "Smoke test": "Bộ test nhanh cho login, claim plot, order, purchase, save/reconnect và admin health.",
    "Regression test": "Chạy suite regression cho gameplay, data integrity, permission và analytics trên release candidate.",
    "Gameplay test": "Playthrough các branch, manual/automation, co-op, season và prestige bằng build release candidate.",
    "Economy test": "Mô phỏng và playtest source/sink, pricing, revenue curve, pacing và correction workflow.",
    "Exploit test": "Chạy toàn bộ exploit matrix, concurrency test và negative permission cases.",
    "Restart test": "Inject restart ở các transaction/state boundary và đối soát sau recovery.",
    "Rollback test": "Diễn tập plugin, DB, content, pack rollback và player-data reconciliation.",
    "Resource-pack failure test": "Inject reject, timeout, hash mismatch, CDN failure và version rollback.",
    "Backup/restore test": "Restore backup vào môi trường cô lập và kiểm tra RPO/data/gameplay.",
    "Moderation test": "Kiểm tra report, sanction, audit, escalation, privacy và permission boundary.",
    "Release acceptance test": "Chạy gate cuối trên artifact, config, migration và pack đúng version dự kiến phát hành.",
}
DESCRIPTIONS.update(TEST_DESCRIPTIONS)

SPECIFIC_AC = {
    "Mục tiêu Closed Beta": "Tài liệu được Product, Engineering và Operations phê duyệt; nêu ít nhất mục tiêu, non-goal, metric, owner và thời điểm review; không còn placeholder ngoài quyết định được ghi owner/deadline.",
    "Số lượng người chơi thử nghiệm": "Các ô [ĐIỀN] về tester, expected concurrency, peak concurrency và cohort đều có owner/deadline; kế hoạch tuyển không vượt capacity benchmark đã pass.",
    "Thời gian beta": "Có start/end, checkpoint và maintenance window được duyệt; calendar và communication template khớp nhau; ít nhất một tabletop schedule review pass.",
    "Tiêu chí chọn tester": "Checklist áp dụng thử cho ít nhất [ĐIỀN: số hồ sơ mẫu] cho kết quả nhất quán; tiêu chí không chứa dữ liệu không cần thiết và có cách xử lý consent.",
    "Nội dung được phép phát hành": "Manifest ghi chính xác build/config/content/pack version; smoke test xác nhận chỉ nội dung trong allowlist có thể truy cập.",
    "Nội dung bị loại khỏi beta": "Mọi mục bị loại có feature flag hoặc không nằm trong artifact; negative test chứng minh player thường không truy cập được.",
    "Exit criteria": "Mọi tiêu chí có metric, ngưỡng [ĐIỀN], query/test evidence và approver; không được exit khi còn P0 hoặc Release Blocker mở.",
    "Theme thứ hai hoặc prestige route": "Quyết định phương án được duyệt; playthrough từ prerequisite đến completion pass; save/restart/rollback không mất hoặc nhân đôi progress; analytics phân biệt content version.",
    "Build stage mới": "Mua stage một lần trừ tiền và unlock đúng một lần; world projection hoàn tất/reconcile sau restart; rollback không làm plot kẹt.",
    "Purchase node mới": "Test insufficient funds, duplicate click, concurrent click và restart đều pass; một idempotency key tạo tối đa một charge và một unlock.",
    "Recipe mới": "Mỗi station transition đúng state machine; sai station/order bị từ chối; disconnect/restart không clone dish; settlement tối đa một lần.",
    "Customer type mới": "Customer hoàn thành hoặc recovery trong timeout [ĐIỀN]; không giữ table reservation vô hạn; payout nằm trong rule đã duyệt và được ghi analytics.",
    "Staff/automation upgrade": "Automation không tạo order/payout trùng; throughput không vượt config; tắt/restart giữa job khôi phục đúng và manual path vẫn dùng được.",
    "Cosmetic reward": "Reward chỉ thay đổi visual/sound/title; automated comparison xác nhận không đổi revenue, speed, capacity, unlock hoặc season progress rate.",
    "Visual và sound polish": "Checklist trạng thái success/error/pending/degraded pass ở client matrix; không có missing model/âm thanh bắt buộc và fallback vẫn truyền đạt trạng thái.",
    "Early-game pacing": "Dashboard có sample size và cohort filter; target [ĐIỀN] được Product duyệt; playtest/analytics release candidate nằm trong tolerance [ĐIỀN].",
    "Mid-game pacing": "Không có purchase gap vượt ngưỡng [ĐIỀN] ngoài chủ ý; funnel và time-to-node được so với Alpha; mọi thay đổi config có before/after evidence.",
    "First prestige pacing": "Time-to-first-prestige p50/p90 theo cohort nằm trong target [ĐIỀN]; outlier được phân loại; không dùng admin/test account.",
    "Purchase price": "100% node có price, prerequisite và config version; simulation không có node âm/zero ngoài allowlist; playtest xác nhận không deadlock progression.",
    "Revenue curve": "Curve được kiểm tra ở mọi stage/branch; không có giảm ngoài thiết kế hoặc growth spike vượt ngưỡng [ĐIỀN]; kết quả khớp ledger sample.",
    "Manual/automation balance": "Cùng stage và cửa sổ [ĐIỀN], uplift manual nằm trong range [ĐIỀN]; co-op/season không vượt cap tổng được duyệt.",
    "Currency source và sink": "100% transaction type map vào source/sink và ledger reason; tổng aggregate chênh ledger không quá [ĐIỀN]; unknown reason = 0.",
    "Inflation monitoring": "Dashboard hiển thị distribution và flow theo cohort/version; synthetic threshold breach tạo đúng alert; runbook nêu owner và hành động.",
    "Economy rollback plan": "Staging drill revert config và áp dụng correction trên dataset mẫu pass; mọi correction có dry-run, approval, ledger entry và reconciliation report.",
    "Progression bottleneck": "Top bottleneck có metric và reproduction; mỗi thay đổi có hypothesis; release candidate không còn blocker ngoài ngưỡng [ĐIỀN].",
    "Hai progression branch": "Mỗi branch hoàn thành được từ save mới; không dead end; switch/commit rule rõ; analytics ghi branch; reward/value chênh trong tolerance [ĐIỀN].",
    "Unlock timing": "Event order claim→purchase→unlock hợp lệ; p50/p90 nằm target [ĐIỀN]; restart không đổi timestamp hoặc phát lại unlock.",
    "Tutorial completion": "Tất cả bước có event start/success/fail; completion đạt target [ĐIỀN]; abandon/reconnect tiếp tục tại state hợp lệ.",
    "Session pacing": "Trong session bands [ĐIỀN], số meaningful event và wait gap nằm target [ĐIỀN]; đo được bằng query đã review.",
    "Retention objective": "Cohort, return event, timezone và exclusions được chốt; dashboard tái lập được; target [ĐIỀN] có owner nhưng không là lý do bỏ qua blocker chất lượng.",
    "Prestige conversion": "Eligible và converted event đối soát được; tỷ lệ target [ĐIỀN]; prestige test không mất entitlement ngoài rule và không double grant.",
    "Catch-up rule nếu có": "Quyết định có/không được ghi; nếu có, test cap/eligibility/expiry pass và không tăng power trả phí; nếu không, mọi catch-up flag bị tắt.",
    "Season theme": "Theme, start/end, timezone, version và content boundary được duyệt; client fallback vẫn hiểu objective.",
    "Objective": "Progress chỉ tăng từ allowlisted server event; boundary/duplicate/replay tests pass; completion chỉ xảy ra đúng ngưỡng [ĐIỀN].",
    "Điều kiện tham gia": "Test eligible, ineligible, late join, reconnect và permission cases đều cho kết quả mong đợi; reason từ chối hiển thị rõ.",
    "Progress tracking": "Mỗi event có unique/idempotency key; replay [ĐIỀN: số] event không đổi tổng; UI khớp durable state sau reconnect/restart.",
    "Reward": "Claim concurrent/retry cấp tối đa một reward; reward inventory đầy được recovery; kiểm tra stat trước/sau xác nhận không có gameplay/economy advantage.",
    "Anti-AFK": "Idle và macro-like repeated test không tăng progress sau rule threshold [ĐIỀN]; gameplay hợp lệ không bị false positive vượt [ĐIỀN].",
    "Anti-pay-to-win": "Review tất cả purchase/entitlement path cho kết quả 0 đường tăng power/progress; sign-off của Product và moderation/policy owner.",
    "Reset khi hết season": "Boundary test trước/đúng/sau end time pass; progress frozen/archived đúng; claim late tuân policy; không reset permanent data.",
    "Recovery khi restart": "Restart ở trước/sau progress commit và reward claim không làm mất/double; reconciliation report có 0 mismatch chưa giải quyết.",
    "Seasonal analytics": "Event schema validation pass; participation→completion→claim funnel đối soát sample với DB; version/timezone có mặt 100%.",
    "Duplicate item": "Drop, death, hopper/container và reconnect tests không tạo thêm entitlement; mỗi order có tối đa một dish claimable/consumed.",
    "Duplicate payout": "Retry, double event và replay settlement [ĐIỀN: số lần] vẫn chỉ có một ledger payout theo order/idempotency key.",
    "Inventory click/drag": "Mọi click/drag mode trong test matrix không di chuyển, clone hoặc biến server-issued item thành item hợp lệ khác; invariant scan = 0 violation.",
    "Shift-click": "Shift-click giữa player/container/GUI bị chặn hoặc xử lý đúng; entitlement count trước/sau không đổi ngoài transition hợp lệ.",
    "Number-key swap": "Hotbar number-key swap ở mọi protected GUI không bypass validation; không clone/mất entitlement; GUI state refresh đúng.",
    "Double-click": "Double-click collect-to-cursor không gom protected item trái phép; entitlement và inventory reconciliation khớp.",
    "Disconnect trong transaction": "Disconnect tại mỗi transaction boundary cho kết quả committed đúng một lần hoặc chưa commit; reconnect không ở state unknown.",
    "Restart trong transaction": "Forced restart tại mỗi transaction boundary; recovery hoàn tất trong [ĐIỀN: thời gian]; 0 double charge/payout/unlock.",
    "Co-op reward duplication": "Hai thành viên trigger/claim đồng thời vẫn tuân ownership và cap; ledger không có duplicate reward key.",
    "Visitor permission bypass": "Visitor thử toàn bộ purchase/reset/station/reward/admin interaction đều bị từ chối server-side và có audit khi cần.",
    "Purchase/reset race condition": "Concurrent purchase và reset ở [ĐIỀN: số vòng] không gây âm tiền, orphan unlock hoặc sai plot state; fence token loại operation cũ.",
    "Prestige duplication": "Concurrent/retry/restart prestige chỉ tăng prestige count và grant bonus đúng một lần; old plot state được reset/reconcile đúng.",
    "Server restart": "Restart matrix qua order/purchase/build/prestige/season pass; sau startup reconciliation 0 invariant violation và player tiếp tục được.",
    "Plugin reload policy": "Runbook nêu rõ reload không hỗ trợ hoặc allowlist; operator command/test tạo thông báo đúng; không có tài liệu nào hướng dẫn hot reload trái policy.",
    "Database outage": "Failure injection xác nhận write kinh tế bị fail-closed/queued theo thiết kế, main thread không treo, không mất/double; recovery reconciliation pass.",
    "Database latency": "Inject latency [ĐIỀN]; tick threshold vẫn đạt; timeout/backpressure kích hoạt; player nhận trạng thái rõ; pool không cạn kéo dài.",
    "Crash recovery": "Kill process ở workload đại diện; startup recovery hoàn thành trong RTO [ĐIỀN]; committed data giữ nguyên và world projection khớp.",
    "Plot reconciliation": "Corrupt/missing/extra block-entity fixtures được phát hiện; chạy reconcile hai lần cho cùng kết quả; 0 thay đổi durable state ngoài rule.",
    "Entity cleanup": "Sau deactivate/reset/restart, orphan entity/task = 0 theo scan; entity budget trở về baseline trong [ĐIỀN].",
    "Stuck customer recovery": "Inject path failure tại mỗi state; watchdog giải phóng table/queue trong timeout [ĐIỀN]; không trừ tiền/reward player sai.",
    "Pending transaction recovery": "Fixtures pending/committed/unknown được phân loại đúng; rerun recovery idempotent; 0 transaction không có final disposition.",
    "Degraded mode": "Mỗi dependency failure vào đúng mode; action nguy hiểm bị chặn; health/UX/alert đúng; thoát mode không replay transaction trùng.",
    "Plugin rollback": "Staging rollback về build N-1 trong RTO [ĐIỀN]; startup/smoke/reconciliation pass; config/schema compatibility được xác nhận.",
    "Database migration rollback": "Migration rehearsal trên bản sao dữ liệu pass cả rollback hoặc forward-fix; row count/checksum/invariants khớp; backup restore point tồn tại.",
    "Content rollback": "Tắt/revert content version không làm save load lỗi; player có progress liên quan được map/compensate; smoke test pass.",
    "Resource-pack rollback": "Server và client tải đúng previous pack/hash; không download loop; model critical có fallback; compatibility matrix pass.",
    "Economy correction": "Dry-run nêu số account/amount/reason; execution tạo một correction mỗi account; rerun không cộng trùng; audit và communication khớp.",
    "Player-data reconciliation": "Report kiểm tra balance, unlock, prestige, entitlement, season và plot; mismatch = 0 hoặc có ticket/owner được phê duyệt trước reopen.",
    "Rollback verification": "100% checklist item có evidence/timestamp/operator; release chỉ reopen khi smoke, invariants, monitoring và data reconciliation pass.",
    "Communication khi rollback": "Tabletop gửi đúng template cho từng severity/audience trong SLA [ĐIỀN]; message nêu impact, action, ETA/update time và correction.",
    "Download flow": "Accept path tải và activate đúng pack/hash trên client matrix; reconnect dùng cache hợp lệ; telemetry ghi outcome mà không thu dữ liệu thừa.",
    "Versioning": "Mỗi artifact có version/hash bất biến; build→pack matrix không mơ hồ; mismatch test bị chặn hoặc fallback đúng.",
    "Hash validation": "Pack đúng hash tải thành công; artifact bị sửa một byte thất bại có kiểm soát; server log/metric ghi version và reason.",
    "Người chơi từ chối pack": "Reject test tuân policy [ĐIỀN]; không vào trạng thái gameplay nửa vời; message hướng dẫn rõ và moderation không cần can thiệp thủ công.",
    "Download failure": "Timeout/HTTP error/corrupt download đều có retry/fallback đúng; không loop vô hạn; support code xuất hiện trong log/player message.",
    "CDN failure": "Failure drill kích hoạt fallback/maintenance trong [ĐIỀN]; alert đến đúng owner; không phục vụ hash/version sai.",
    "Fallback presentation": "Không pack, mọi action critical vẫn phân biệt được bằng vanilla material/text; playtest hoàn tất core loop không cần đoán custom model.",
    "Client compatibility": "100% client trong support matrix pass login/download/core loop; ngoài matrix nhận thông báo unsupported rõ ràng.",
    "Rollback pack version": "Chuyển N→N-1→N trong test không cache sai hoặc loop; hash/version telemetry đúng; core visual smoke pass.",
    "Co-op permission": "Matrix owner/member cho mọi action pass; revoke có hiệu lực trong [ĐIỀN]; concurrent action không vượt ownership/idempotency rule.",
    "Visitor permission": "Matrix visitor pass 100%; purchase/reset/reward/stock/admin action bị deny server-side; allowlisted social action vẫn hoạt động.",
    "Grief prevention": "Automated/manual matrix cho block, container, entity, fluid, piston, explosion và vehicle có 0 unauthorized world change.",
    "Chat moderation": "Test allow/block/mute/appeal/audit pass theo policy; moderator action có reason; false-positive samples được review trước beta.",
    "Report flow": "Report hợp lệ tạo ID và audit record; spam/rate-limit/invalid target xử lý đúng; moderator xem đủ context theo least privilege.",
    "Player sanction": "Mỗi sanction type có apply/expire/revoke/appeal test; enforcement qua reconnect/restart; audit log đầy đủ actor/reason/duration.",
    "Audit log": "100% action nhạy cảm trong allowlist tạo log đủ field; tamper/access test pass; query theo correlation ID trả đúng chuỗi sự kiện.",
    "Admin escalation": "Tabletop SEV [ĐIỀN] đến đúng on-call trong SLA; break-glass được audit và revoke; escalation contacts đã xác nhận.",
    "Privacy rule": "Role test chứng minh chỉ người có quyền xem field cần thiết; export/delete/retention path được diễn tập; log không chứa secret hoặc chat ngoài policy.",
    "Deployment checklist": "Một dry run release candidate hoàn tất 100% mục bắt buộc, có owner/evidence; thiếu mục P0 tự động dừng deploy.",
    "Canary release": "Canary nhận đúng version; gate theo [ĐIỀN: metrics/thresholds] quyết định promote/rollback; rollback drill nằm trong RTO.",
    "Monitoring dashboard": "Synthetic gameplay và failure injection xuất hiện đúng dashboard trong [ĐIỀN: delay]; không có panel critical No Data không được giải thích.",
    "Alert rule": "Mỗi alert có test fire/resolution, severity, runbook và owner; signal dưới ngưỡng không page; signal vượt ngưỡng page trong [ĐIỀN].",
    "Incident severity": "[ĐIỀN: số tình huống] tabletop được phân loại nhất quán theo matrix; mỗi severity có response/communication target và authority rõ.",
    "Incident response": "Tabletop từ alert đến resolved/postmortem pass; timeline/correlation/evidence đầy đủ; containment không làm hỏng dữ liệu.",
    "Backup schedule": "Backup chạy theo RPO [ĐIỀN], encrypted và off-host; job failure alert; restore sampling xác nhận artifact đọc được.",
    "Restore procedure": "Restore rehearsal đạt RTO [ĐIỀN] và RPO [ĐIỀN]; schema/checksum/invariants/gameplay smoke pass trước cutover.",
    "Maintenance mode": "Player thường không tạo write mới; admin health/restore access hoạt động; message/ETA đúng; thoát mode không gây login storm ngoài budget.",
    "Player communication": "Templates được duyệt; tabletop gửi update theo cadence [ĐIỀN]; nội dung không hứa thời gian/chức năng chưa xác nhận.",
    "Error tracking": "Inject exception có event trong [ĐIỀN: delay] với build/correlation; secret/PII fixtures bị scrub; duplicate được group đúng.",
    "Performance tracking": "Metrics có p50/p95/p99 và build tag; benchmark query tái lập được; sampling overhead dưới [ĐIỀN].",
    "Economy tracking": "Sample transaction đối soát 100% với ledger; duplicate event không double-count; dashboard filter đúng config/cohort/version.",
    "Progression tracking": "Scripted playthrough tạo đúng ordered funnel events; reconnect/replay không nhân đôi; admin/test account được gắn để loại trừ.",
    "Retention tracking": "Query cohort được peer review và tái lập trên fixture; timezone/boundary tests pass; bot/admin/test account exclusions đúng.",
    "Tester feedback form": "Form submit thành công trên support clients; required fields/rate limit/consent pass; ticket liên kết build và tester pseudonymous ID.",
    "Bug report template": "[ĐIỀN: số báo cáo mẫu] đủ dữ liệu để triage/reproduce; field data-risk và release-blocker bắt buộc; không yêu cầu secret/PII.",
    "Feedback prioritization": "Triage rehearsal cho [ĐIỀN: số item] tạo priority/owner/SLA nhất quán; P0/Release Blocker không nằm ngoài queue xử lý.",
    "Privacy và data retention": "Data inventory được duyệt; TTL/deletion/access test pass cho từng store; không thu field không có purpose/owner.",
    "Expected beta concurrency": "Giả định player/active plot/NPC/entity/path/DB workload được điền và phê duyệt; claim capacity liên kết tới benchmark evidence.",
    "Peak concurrency test": "Ở peak [ĐIỀN], tất cả tick/memory/DB/error threshold [ĐIỀN] pass trong steady-state; report lưu config, hardware và workload.",
    "Active plot count": "Test tăng active plot tới [ĐIỀN]; metric đếm đúng; inactive plot về 0 NPC và không force-load chunk ngoài policy.",
    "NPC count": "Budget [ĐIỀN] không bị vượt trong workload; plot inactive cleanup đúng; overload test không phá tick/data invariants.",
    "Entity tracking": "Tracked entity/client-visible counts nằm budget [ĐIỀN]; reconnect/plot switch không làm count tăng dần; orphan scan = 0.",
    "Pathfinding load": "Request rate/cost nằm threshold [ĐIỀN]; idle customer không repath liên tục; injected stuck path được watchdog recovery.",
    "Tick time": "p50/p95/p99 MSPT ở benchmark/soak dưới ngưỡng [ĐIỀN]; không có main-thread DB/file/HTTP call trong profile.",
    "Memory": "Heap/native/GC nằm ngưỡng [ĐIỀN]; slope sau warm-up không vượt [ĐIỀN]; heap dump/profile không có retained leak blocker.",
    "Database pool": "Peak test không có pool exhaustion kéo dài; acquire/query p95 dưới [ĐIỀN]; timeout không block main thread và recovery pass.",
    "Long soak test": "Chạy đủ [ĐIỀN: duration]; không crash, leak slope, data invariant violation hoặc alert P0/P1 chưa giải quyết; restart sau soak pass.",
    "Smoke test": "100% case P0 login→claim→order→purchase→save/reconnect và admin health pass trên exact release artifact; không có blocker mở.",
    "Regression test": "Toàn bộ suite bắt buộc pass; flaky test có owner và không che P0/P1; report gắn commit/build/config/pack versions.",
    "Gameplay test": "Playthrough mỗi branch và prestige/season/co-op path pass; không softlock; mọi defect được triage theo severity trước sign-off.",
    "Economy test": "Simulation và playtest nằm pacing/curve/source-sink thresholds [ĐIỀN]; ledger reconciliation pass; rollback/correction rehearsal pass.",
    "Exploit test": "100% exploit cases trong nhóm Exploit Testing pass; 0 duplication, negative balance, unauthorized action hoặc unresolved invariant violation.",
    "Restart test": "Restart ở mọi checkpoint bắt buộc pass; RTO [ĐIỀN]; 0 mất/double transaction/progress và reconciliation sạch.",
    "Rollback test": "Plugin/DB/content/pack rollback đạt RTO/RPO [ĐIỀN]; smoke + reconciliation + monitoring pass trước reopen.",
    "Resource-pack failure test": "Reject, timeout, corrupt/hash mismatch, CDN failure và rollback version cho đúng policy/fallback; không download loop.",
    "Backup/restore test": "Restore bản backup beta gần nhất vào môi trường cô lập đạt RPO/RTO [ĐIỀN]; checksum/invariants/smoke pass.",
    "Moderation test": "Report→triage→sanction→appeal/audit/escalation path pass; permission/privacy negative tests có 0 unauthorized access.",
    "Release acceptance test": "Tất cả Release Blocker hoàn thành; Beta Exit Criteria được ký; exact artifact/config/schema/pack pass smoke, regression và operational readiness review.",
}

EXPLOIT_ITEMS = {item for group, items in GROUPS if group == "Exploit Testing" for item in items}
BLOCKERS = {
    "Exit criteria", "Economy rollback plan", "Anti-pay-to-win", "Recovery khi restart",
    *EXPLOIT_ITEMS,
    "Server restart", "Database outage", "Crash recovery", "Plot reconciliation",
    "Pending transaction recovery", "Degraded mode", "Plugin rollback", "Database migration rollback",
    "Economy correction", "Player-data reconciliation", "Rollback verification", "Hash validation",
    "Download failure", "CDN failure", "Fallback presentation", "Rollback pack version",
    "Co-op permission", "Visitor permission", "Grief prevention", "Audit log", "Deployment checklist",
    "Monitoring dashboard", "Alert rule", "Incident response", "Backup schedule", "Restore procedure",
    "Error tracking", "Performance tracking", "Economy tracking", "Privacy và data retention",
    "Peak concurrency test", "Tick time", "Memory", "Database pool", "Long soak test",
    "Smoke test", "Regression test", "Exploit test", "Restart test", "Rollback test",
    "Resource-pack failure test", "Backup/restore test", "Release acceptance test",
}

P1_ITEMS = {
    "Mục tiêu Closed Beta", "Số lượng người chơi thử nghiệm", "Thời gian beta", "Tiêu chí chọn tester",
    "Nội dung được phép phát hành", "Nội dung bị loại khỏi beta", "Early-game pacing", "Mid-game pacing",
    "First prestige pacing", "Purchase price", "Revenue curve", "Manual/automation balance",
    "Currency source và sink", "Inflation monitoring", "Progression bottleneck", "Hai progression branch",
    "Unlock timing", "Tutorial completion", "Session pacing", "Prestige conversion", "Objective",
    "Progress tracking", "Anti-AFK", "Reset khi hết season", "Seasonal analytics", "Database latency",
    "Entity cleanup", "Stuck customer recovery", "Content rollback", "Resource-pack rollback",
    "Communication khi rollback", "Download flow", "Versioning", "Người chơi từ chối pack",
    "Client compatibility", "Chat moderation", "Report flow", "Player sanction", "Admin escalation",
    "Privacy rule", "Canary release", "Incident severity", "Maintenance mode", "Player communication",
    "Progression tracking", "Retention tracking", "Feedback prioritization", "Expected beta concurrency",
    "Active plot count", "NPC count", "Entity tracking", "Pathfinding load", "Gameplay test",
    "Economy test", "Moderation test",
}

COMPLEXITY = {
    "XL": {"Theme thứ hai hoặc prestige route", "Database migration rollback", "Long soak test"},
    "L": {
        "Build stage mới", "Recipe mới", "Staff/automation upgrade", "Revenue curve", "Hai progression branch",
        "Progress tracking", "Anti-AFK", "Recovery khi restart", "Purchase/reset race condition",
        "Server restart", "Database outage", "Crash recovery", "Plot reconciliation",
        "Pending transaction recovery", "Degraded mode", "Plugin rollback", "Content rollback",
        "Economy correction", "Player-data reconciliation", "CDN failure", "Fallback presentation",
        "Co-op permission", "Grief prevention", "Audit log", "Incident response", "Restore procedure",
        "Performance tracking", "Economy tracking", "Peak concurrency test", "Pathfinding load",
        "Tick time", "Memory", "Database pool", "Regression test", "Exploit test", "Restart test",
        "Rollback test", "Backup/restore test", "Release acceptance test",
    },
    "S": {
        "Mục tiêu Closed Beta", "Số lượng người chơi thử nghiệm", "Thời gian beta", "Tiêu chí chọn tester",
        "Nội dung được phép phát hành", "Nội dung bị loại khỏi beta", "Exit criteria", "Cosmetic reward",
        "Season theme", "Điều kiện tham gia", "Reward", "Plugin reload policy", "Communication khi rollback",
        "Versioning", "Người chơi từ chối pack", "Client compatibility", "Incident severity",
        "Backup schedule", "Player communication", "Tester feedback form", "Bug report template",
        "Feedback prioritization", "Expected beta concurrency", "Smoke test",
    },
    "XS": {"Privacy rule"},
}

DEPS_BY_GROUP = {
    "Beta Scope": "MVP Alpha findings; Product sign-off",
    "Content Expansion": "Beta Scope; content pipeline; versioned config",
    "Economy Tuning": "Alpha analytics; Economy tracking; Product targets [ĐIỀN]",
    "Progression Tuning": "Alpha progression analytics; Economy Tuning",
    "Seasonal Objectives": "Beta Scope; Analytics; server-authoritative event model",
    "Exploit Testing": "Stable release candidate; transaction/permission invariants",
    "Reliability": "Persistence model; observability; recovery jobs",
    "Rollback": "Versioned artifact/config/schema; backups; runbook",
    "Resource Pack": "Pack artifact; CDN; client support matrix [ĐIỀN]",
    "Multiplayer và Moderation": "Permission model; policy; audit storage",
    "Operations": "Release candidate; infrastructure; on-call ownership",
    "Analytics và Feedback": "Event schema; privacy review; dashboard pipeline",
    "Performance": "Production-like hardware; workload model [ĐIỀN]",
    "Beta Testing": "Release candidate; test environment; preceding workstreams",
}

PLAYER_VALUE_BY_GROUP = {
    "Beta Scope": "Kỳ vọng beta rõ ràng, trải nghiệm tập trung và ít thay đổi ngoài dự kiến.",
    "Content Expansion": "Có mục tiêu mới đáng khám phá mà không đánh đổi độ ổn định.",
    "Economy Tuning": "Nhịp kiếm và tiêu tiền công bằng, ít grind hoặc spike bất hợp lý.",
    "Progression Tuning": "Luôn thấy mục tiêu tiếp theo và có lựa chọn tiến triển có ý nghĩa.",
    "Seasonal Objectives": "Có lý do quay lại theo mùa nhưng không pay-to-win.",
    "Exploit Testing": "Tiến trình và economy công bằng, không bị phá bởi duplication/bypass.",
    "Reliability": "Ít mất tiến trình, downtime và trạng thái kẹt khi hệ thống lỗi.",
    "Rollback": "Sự cố được phục hồi nhanh, minh bạch và hạn chế mất dữ liệu.",
    "Resource Pack": "Visual nhất quán; vẫn chơi và hiểu trạng thái khi pack gặp lỗi.",
    "Multiplayer và Moderation": "Co-op/social an toàn, không grief và có hỗ trợ khi bị lạm dụng.",
    "Operations": "Server ổn định, sự cố được phát hiện và xử lý nhanh.",
    "Analytics và Feedback": "Phản hồi dẫn đến sửa lỗi đúng ưu tiên mà tôn trọng privacy.",
    "Performance": "Gameplay mượt và ổn định trong tải beta thực tế.",
    "Beta Testing": "Release candidate đáng tin cậy trước khi mời tester thật.",
}


def priority(item):
    if item in BLOCKERS:
        return "P0"
    if item in P1_ITEMS:
        return "P1"
    if item in {"Cosmetic reward", "Visual và sound polish", "Season theme", "Reward", "Catch-up rule nếu có"}:
        return "P3"
    return "P2"


def complexity(item):
    for size, items in COMPLEXITY.items():
        if item in items:
            return size
    return "M"


def description(group, item):
    return DESCRIPTIONS.get(item, f"Thiết kế, triển khai và xác minh hạng mục {item} cho Closed Beta theo phạm vi đã duyệt.")


def acceptance(group, item):
    if item in SPECIFIC_AC:
        return SPECIFIC_AC[item]
    return f"Test plan cho {item} được review; happy path, boundary, permission, reconnect và restart cases liên quan đều pass; 0 lỗi P0/P1 mở và có evidence gắn release candidate."


def setup_styles(ws):
    navy = "17324D"
    blue = "1F4E78"
    light_blue = "D9EAF7"
    pale_blue = "EAF3F8"
    red = "C00000"
    pale_red = "FCE4D6"
    yellow = "FFF2CC"
    green = "E2F0D9"
    white = "FFFFFF"
    thin = Side(style="thin", color="B7C9D6")

    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A7"
    ws.auto_filter.ref = f"A6:M{6 + sum(len(items) for _, items in GROUPS)}"
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 26
    ws.merge_cells("A1:M1")
    ws["A1"] = "MINECRAFT RESTAURANT TYCOON | PHASE 3 - CLOSED BETA"
    ws["A1"].font = Font(name="Arial", size=16, bold=True, color=white)
    ws["A1"].fill = PatternFill("solid", fgColor=navy)
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.merge_cells("A2:M2")
    ws["A2"] = "Ưu tiên sửa lỗi, ổn định và khả năng phục hồi. Ô [ĐIỀN: ...] là quyết định sản phẩm/vận hành chưa được tự quyết định."
    ws["A2"].font = Font(name="Arial", size=10, italic=True, color="334E68")
    ws["A2"].fill = PatternFill("solid", fgColor=light_blue)
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")

    legends = [
        ("A3", "Hướng dẫn", blue, white),
        ("B3", "Chỉ sửa placeholder [ĐIỀN: ...], Người phụ trách, Trạng thái, Tiến độ % và Ghi chú quyết định.", pale_blue, "1F1F1F"),
        ("A4", "Ví dụ định dạng", blue, white),
        ("B4", "[ĐIỀN: target] | Owner: [ĐIỀN: tên/role] | Trạng thái: Đang thiết kế | Tiến độ: 25", yellow, "1F1F1F"),
        ("A5", "Release blocker", red, white),
        ("B5", "Dòng nền đỏ, Priority P0 và Ghi chú bắt đầu bằng RELEASE BLOCKER. Không phát hành khi còn blocker chưa Hoàn thành.", pale_red, red),
    ]
    for cell, value, fill, color in legends:
        ws[cell] = value
        ws[cell].font = Font(name="Arial", size=9, bold=cell.startswith("A"), color=color)
        ws[cell].fill = PatternFill("solid", fgColor=fill)
        ws[cell].alignment = Alignment(wrap_text=True, vertical="center")
    for row in range(3, 6):
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=13)
        ws.row_dimensions[row].height = 28

    headers = [
        "ID", "Nhóm công việc", "Tính năng/Hạng mục", "Mô tả cần điền", "Player Value", "Phụ thuộc",
        "Độ ưu tiên", "Độ phức tạp", "Người phụ trách", "Trạng thái", "Tiến độ %",
        "Acceptance Criteria", "Ghi chú",
    ]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(6, col, header)
        cell.font = Font(name="Arial", size=10, bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=blue)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    ws.row_dimensions[6].height = 36

    widths = [16, 24, 30, 54, 38, 35, 13, 14, 22, 18, 12, 72, 28]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    return {
        "thin": thin, "pale_blue": pale_blue, "pale_red": pale_red, "yellow": yellow,
        "green": green, "red": red, "blue": blue, "navy": navy, "white": white,
    }


def add_validations(ws, start_row, end_row):
    status = DataValidation(
        type="list",
        formula1='"Chưa làm,Đang thiết kế,Sẵn sàng,Đang làm,Đang kiểm thử,Bị chặn,Hoàn thành,Hủy"',
        allow_blank=False,
    )
    priority_dv = DataValidation(type="list", formula1='"P0,P1,P2,P3"', allow_blank=False)
    complexity_dv = DataValidation(type="list", formula1='"XS,S,M,L,XL"', allow_blank=False)
    progress = DataValidation(type="whole", operator="between", formula1="0", formula2="100", allow_blank=False)
    progress.error = "Tiến độ phải là số nguyên từ 0 đến 100."
    progress.errorTitle = "Tiến độ không hợp lệ"
    for dv in (status, priority_dv, complexity_dv, progress):
        ws.add_data_validation(dv)
    status.add(f"J{start_row}:J{end_row}")
    priority_dv.add(f"G{start_row}:G{end_row}")
    complexity_dv.add(f"H{start_row}:H{end_row}")
    progress.add(f"K{start_row}:K{end_row}")


def add_summary_table(ws, start_row, title, headers, rows, styles):
    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=13)
    title_cell = ws.cell(start_row, 1, title)
    title_cell.font = Font(name="Arial", size=12, bold=True, color=styles["white"])
    title_cell.fill = PatternFill("solid", fgColor=styles["navy"])
    title_cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[start_row].height = 24

    header_row = start_row + 1
    for col, header in enumerate(headers, 1):
        cell = ws.cell(header_row, col, header)
        cell.font = Font(name="Arial", size=9, bold=True, color=styles["white"])
        cell.fill = PatternFill("solid", fgColor=styles["blue"])
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(top=styles["thin"], bottom=styles["thin"], left=styles["thin"], right=styles["thin"])
    for r_index, row in enumerate(rows, header_row + 1):
        for c_index, value in enumerate(row, 1):
            cell = ws.cell(r_index, c_index, value)
            cell.font = Font(name="Arial", size=9, color="1F1F1F")
            cell.fill = PatternFill("solid", fgColor="FFFFFF" if r_index % 2 else styles["pale_blue"])
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(top=styles["thin"], bottom=styles["thin"], left=styles["thin"], right=styles["thin"])
        ws.row_dimensions[r_index].height = 48
    return header_row + len(rows) + 2


def build_workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = "Phase 3 - Closed Beta"
    styles = setup_styles(ws)
    row = 7
    group_number = 0
    for group, items in GROUPS:
        group_number += 1
        for item_number, item in enumerate(items, 1):
            blocker = item in BLOCKERS
            values = [
                f"P3-{group_number:02d}-{item_number:02d}", group, item, description(group, item),
                PLAYER_VALUE_BY_GROUP[group], DEPS_BY_GROUP[group], priority(item), complexity(item),
                "[ĐIỀN: tên/role]", "Chưa làm", 0, acceptance(group, item),
                "RELEASE BLOCKER | Không phát hành nếu chưa Hoàn thành." if blocker else "[ĐIỀN: quyết định, link evidence hoặc issue]",
            ]
            for col, value in enumerate(values, 1):
                cell = ws.cell(row, col, value)
                cell.font = Font(name="Arial", size=9, color=styles["red"] if blocker and col in (1, 7, 13) else "1F1F1F", bold=blocker and col in (1, 7, 13))
                fill = styles["pale_red"] if blocker else ("FFFFFF" if group_number % 2 else styles["pale_blue"])
                if col in (4, 9, 13) and "[ĐIỀN:" in str(value) and not blocker:
                    fill = styles["yellow"]
                cell.fill = PatternFill("solid", fgColor=fill)
                cell.alignment = Alignment(
                    horizontal="center" if col in (1, 7, 8, 10, 11) else "left",
                    vertical="top", wrap_text=True,
                )
                cell.border = Border(top=styles["thin"], bottom=styles["thin"], left=styles["thin"], right=styles["thin"])
            ws.row_dimensions[row].height = 78
            row += 1

    plan_end_row = row - 1
    add_validations(ws, 7, plan_end_row)
    ws.conditional_formatting.add(
        f"A7:M{plan_end_row}",
        FormulaRule(formula=["$J7=\"Hoàn thành\""], fill=PatternFill("solid", fgColor=styles["green"])),
    )
    ws.conditional_formatting.add(
        f"A7:M{plan_end_row}",
        FormulaRule(formula=["$J7=\"Bị chặn\""], fill=PatternFill("solid", fgColor=styles["pale_red"])),
    )

    row += 2
    milestone_rows = [
        ("M1", "Scope & telemetry lock", "[ĐIỀN: ngày]", "Beta scope, metric definitions, event schema và product placeholders quan trọng được duyệt.", "[ĐIỀN: owner]", "Chưa làm"),
        ("M2", "Content/economy candidate", "[ĐIỀN: ngày]", "Content expansion đã chọn; economy/progression config candidate có evidence từ Alpha và playtest.", "[ĐIỀN: owner]", "Chưa làm"),
        ("M3", "Reliability & exploit gate", "[ĐIỀN: ngày]", "Exploit, restart, recovery, rollback, resource-pack failure và backup/restore drills pass.", "[ĐIỀN: owner]", "Chưa làm"),
        ("M4", "Operational readiness", "[ĐIỀN: ngày]", "Runbook, monitoring, alerts, moderation, privacy, escalation và communication được tabletop.", "[ĐIỀN: owner]", "Chưa làm"),
        ("M5", "Closed Beta release", "[ĐIỀN: ngày]", "Release acceptance pass; 0 blocker mở; approvers ký Beta Exit Criteria để bắt đầu beta.", "[ĐIỀN: owner]", "Chưa làm"),
    ]
    row = add_summary_table(ws, row, "Milestones", ["ID", "Milestone", "Ngày mục tiêu", "Điều kiện hoàn thành", "Owner", "Trạng thái"], milestone_rows, styles)

    risk_rows = [
        ("R-01", "Economy tune từ sample Alpha không đại diện", "[ĐIỀN]", "[ĐIỀN]", "Phân cohort, ghi confidence/sample size, canary config và chuẩn bị economy rollback.", "[ĐIỀN: owner]", "Chưa làm"),
        ("R-02", "Transaction duplication hoặc data inconsistency khi restart", "[ĐIỀN]", "Rất cao", "Idempotency, ledger invariant, crash injection, pending recovery và reconciliation gate.", "[ĐIỀN: owner]", "Chưa làm"),
        ("R-03", "NPC/pathfinding vượt capacity trong tải thật", "[ĐIỀN]", "Cao", "Benchmark theo active plot/NPC/path request; budget, inactive simulation và degraded mode.", "[ĐIỀN: owner]", "Chưa làm"),
        ("R-04", "Resource pack/CDN failure chặn tester", "[ĐIỀN]", "Cao", "Hash/version pin, fallback presentation, retry, CDN drill và rollback pack.", "[ĐIỀN: owner]", "Chưa làm"),
        ("R-05", "Seasonal objective tạo AFK/P2W hoặc exploit", "[ĐIỀN]", "Cao", "Server-authoritative progress, anti-AFK, P2W review và reward cosmetic-only.", "[ĐIỀN: owner]", "Chưa làm"),
        ("R-06", "Moderator/on-call chưa sẵn sàng", "[ĐIỀN]", "Cao", "Permission/audit review, incident tabletop, escalation roster và communication templates.", "[ĐIỀN: owner]", "Chưa làm"),
        ("R-07", "Thêm content làm trễ hardening", "[ĐIỀN]", "Trung bình", "Timebox content, feature flag; cắt content trước reliability/release blocker.", "[ĐIỀN: owner]", "Chưa làm"),
    ]
    row = add_summary_table(ws, row, "Risks", ["ID", "Risk", "Khả năng", "Ảnh hưởng", "Mitigation/Contingency", "Owner", "Trạng thái"], risk_rows, styles)

    blocker_rows = [
        ("RB-01", "Data integrity & exploit", "Tất cả exploit matrix pass; 0 double payout/charge/unlock/reward; invariant reconciliation sạch.", "P3-06-*", "[ĐIỀN: owner]", "Chưa làm"),
        ("RB-02", "Restart/crash/degraded recovery", "Restart/crash/DB failure drills đạt RTO [ĐIỀN], không mất/nhân đôi data và không block main thread.", "P3-07-*", "[ĐIỀN: owner]", "Chưa làm"),
        ("RB-03", "Rollback readiness", "Plugin, DB, content, pack và economy rollback diễn tập pass; player-data reconciliation = 0 mismatch mở.", "P3-08-*", "[ĐIỀN: owner]", "Chưa làm"),
        ("RB-04", "Resource-pack resilience", "Reject/download/hash/CDN/fallback/version rollback tests pass trên client matrix [ĐIỀN].", "P3-09-*", "[ĐIỀN: owner]", "Chưa làm"),
        ("RB-05", "Permission, grief & audit", "Co-op/visitor/grief negative tests pass; admin/economy actions có audit và privacy controls.", "P3-10-*", "[ĐIỀN: owner]", "Chưa làm"),
        ("RB-06", "Operational readiness", "Deploy/canary/monitor/alert/incident/backup/restore runbook được drill trên release candidate.", "P3-11-*", "[ĐIỀN: owner]", "Chưa làm"),
        ("RB-07", "Performance capacity", "Peak và soak workload [ĐIỀN] pass tick/memory/DB/entity/path thresholds [ĐIỀN].", "P3-13-*", "[ĐIỀN: owner]", "Chưa làm"),
        ("RB-08", "Release acceptance", "Smoke/regression/exploit/restart/rollback/resource-pack/restore gates pass; 0 P0 hoặc blocker mở.", "P3-14-*", "[ĐIỀN: owner]", "Chưa làm"),
    ]
    row = add_summary_table(ws, row, "Release Blockers", ["ID", "Blocker", "Điều kiện gỡ blocker", "Liên kết hạng mục", "Owner", "Trạng thái"], blocker_rows, styles)

    exit_rows = [
        ("BE-01", "Scope", "Beta scope, tester/time window và content allowlist/denylist được phê duyệt; không còn quyết định bắt buộc vô chủ.", "[ĐIỀN: approver]", "Chưa làm"),
        ("BE-02", "Quality", "0 P0, 0 Release Blocker; P1 còn lại có risk acceptance, owner và deadline được duyệt.", "[ĐIỀN: approver]", "Chưa làm"),
        ("BE-03", "Data integrity", "Exploit/restart/crash/rollback tests không tạo duplicate, loss hoặc invariant mismatch chưa xử lý.", "[ĐIỀN: approver]", "Chưa làm"),
        ("BE-04", "Economy & progression", "Các target pacing/economy [ĐIỀN] đạt tolerance [ĐIỀN] trên sample đủ điều kiện; rollback plan pass.", "[ĐIỀN: approver]", "Chưa làm"),
        ("BE-05", "Performance", "Expected/peak workload [ĐIỀN] đạt tick/memory/DB/entity/path và soak thresholds [ĐIỀN].", "[ĐIỀN: approver]", "Chưa làm"),
        ("BE-06", "Resource pack", "Client matrix và reject/failure/CDN/fallback/rollback pack scenarios pass.", "[ĐIỀN: approver]", "Chưa làm"),
        ("BE-07", "Operations", "Deployment, monitoring, alerting, incident, backup/restore, rollback và communication drills pass.", "[ĐIỀN: approver]", "Chưa làm"),
        ("BE-08", "Moderation & privacy", "Permission/grief/report/sanction/audit/escalation/privacy tests pass; staff roster sẵn sàng.", "[ĐIỀN: approver]", "Chưa làm"),
        ("BE-09", "Season fairness", "Seasonal objective recovery/anti-AFK/anti-P2W/reward/reset/analytics tests pass.", "[ĐIỀN: approver]", "Chưa làm"),
    ]
    row = add_summary_table(ws, row, "Beta Exit Criteria", ["ID", "Nhóm tiêu chí", "Tiêu chí kiểm thử", "Approver", "Kết quả"], exit_rows, styles)

    dod_rows = [
        ("DoD-01", "Thiết kế & scope", "Acceptance Criteria, dependency, owner, priority và failure behavior được review; placeholder bắt buộc đã có quyết định hoặc owner/deadline."),
        ("DoD-02", "Implementation", "Code/config/content versioned; server-authoritative validation, idempotency và feature flag/rollback path khi phù hợp."),
        ("DoD-03", "Testing", "Happy path, boundary, negative permission, concurrent/retry, reconnect/restart và rollback tests liên quan pass."),
        ("DoD-04", "Data", "Migration/ledger/event schema review; reconciliation và privacy/retention controls pass; không có invariant mismatch mở."),
        ("DoD-05", "Observability", "Metrics/log/error events có build/version/correlation; dashboard và alert test fire/resolution pass."),
        ("DoD-06", "Operations", "Runbook deploy/recover/rollback cập nhật; backup/restore hoặc tabletop/drill evidence đính kèm khi có impact vận hành."),
        ("DoD-07", "Player experience", "Thông báo success/error/degraded rõ; fallback hoạt động; localization/text và accessibility/readability được review."),
        ("DoD-08", "Sign-off", "Không còn P0/P1 không được chấp thuận; Release Blocker liên quan đã gỡ; Product, QA và Operations ký theo ownership."),
    ]
    add_summary_table(ws, row, "Definition of Done", ["ID", "Khía cạnh", "Điều kiện Done"], dod_rows, styles)

    ws.print_title_rows = "1:6"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.auto_filter.ref = f"A6:M{plan_end_row}"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.save(OUTPUT)
    return plan_end_row


def verify(plan_end_row):
    wb = load_workbook(OUTPUT, data_only=False)
    ws = wb["Phase 3 - Closed Beta"]
    expected_items = sum(len(items) for _, items in GROUPS)
    assert plan_end_row - 6 == expected_items
    assert ws.max_column == 13
    assert ws["A6"].value == "ID"
    assert ws["M6"].value == "Ghi chú"
    ids = [ws.cell(r, 1).value for r in range(7, plan_end_row + 1)]
    assert len(ids) == len(set(ids)) == expected_items
    for r in range(7, plan_end_row + 1):
        assert ws.cell(r, 12).value and len(ws.cell(r, 12).value) >= 40
        assert ws.cell(r, 7).value in {"P0", "P1", "P2", "P3"}
        assert ws.cell(r, 8).value in {"XS", "S", "M", "L", "XL"}
        assert ws.cell(r, 10).value in {"Chưa làm", "Đang thiết kế", "Sẵn sàng", "Đang làm", "Đang kiểm thử", "Bị chặn", "Hoàn thành", "Hủy"}
        if ws.cell(r, 7).value == "P0":
            assert str(ws.cell(r, 13).value).startswith("RELEASE BLOCKER")
    assert any(cell.value == "Milestones" for row in ws.iter_rows() for cell in row)
    assert any(cell.value == "Risks" for row in ws.iter_rows() for cell in row)
    assert any(cell.value == "Release Blockers" for row in ws.iter_rows() for cell in row)
    assert any(cell.value == "Beta Exit Criteria" for row in ws.iter_rows() for cell in row)
    assert any(cell.value == "Definition of Done" for row in ws.iter_rows() for cell in row)
    print(f"Created: {OUTPUT}")
    print(f"Plan items: {expected_items}; release blockers: {len(BLOCKERS)}; rows: {ws.max_row}; columns: {ws.max_column}")


if __name__ == "__main__":
    end = build_workbook()
    verify(end)
