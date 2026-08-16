from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


OUTPUT = Path(__file__).with_name("Minecraft_Restaurant_Tycoon_Phase_2_MVP_Alpha.xlsx")

HEADERS = [
    "ID", "Nhóm công việc", "Tính năng/Hạng mục", "Mô tả cần điền",
    "Player Value", "Phụ thuộc", "Độ ưu tiên", "Độ phức tạp",
    "Người phụ trách", "Trạng thái", "Tiến độ %", "Acceptance Criteria", "Ghi chú",
]

STATUS = ["Chưa làm", "Đang thiết kế", "Sẵn sàng", "Đang làm", "Đang kiểm thử", "Bị chặn", "Hoàn thành", "Hủy"]
PRIORITY = ["P0", "P1", "P2", "P3"]
COMPLEXITY = ["XS", "S", "M", "L", "XL"]


def item(item_id, group, feature, description, value, dependency, priority, complexity, criteria, note=""):
    return [item_id, group, feature, description, value, dependency, priority, complexity,
            "[Chọn owner]", "Chưa làm", 0, criteria, note]


rows = [
    # 1. Progression Expansion
    item("PRG-01", "1. Progression Expansion", "40-60 purchase node", "Thiết kế và cấu hình tổng cộng [40-60] node; mỗi node có ID, loại, prerequisite, effect, stage và vị trí control.", "Luôn có mục tiêu mua tiếp theo rõ ràng và nhịp tăng trưởng liên tục.", "Vertical slice; ECO-01; CON-03", "P0", "XL", "Có 40-60 node hợp lệ; 100% ID duy nhất; 100% node reachable từ node đầu; không có cycle ngoài cấu hình được duyệt; validator chạy 0 lỗi."),
    item("PRG-02", "1. Progression Expansion", "Hai progression branch", "Tạo hai nhánh lựa chọn [Nhánh A] và [Nhánh B], có điểm rẽ và điểm hội tụ rõ ràng; không tự điền theme/reward.", "Cho người chơi quyền lựa chọn và tăng replay value.", "PRG-01; ECO-01", "P0", "L", "Cả 2 nhánh đều hoàn thành được độc lập; mỗi nhánh có >= [số node] node riêng; chênh lệch P50 time-to-complete giữa hai nhánh <= [ngưỡng %]; không có soft-lock trong 20 lượt test/nhánh."),
    item("PRG-03", "1. Progression Expansion", "Build stage mở rộng", "Ánh xạ purchase node vào canonical build stage và world operation có thể retry/replay.", "Mỗi lần mua tạo thay đổi vật lý dễ nhận biết trong nhà hàng.", "PRG-01; PER-05", "P0", "XL", "100% major node đổi geometry trong <= [giây] sau commit; retry 3 lần cho cùng operation cho kết quả world giống nhau; stage validator báo 0 block/entity ngoài plot."),
    item("PRG-04", "1. Progression Expansion", "Mid-game pacing", "Tune band giữa game bằng time-to-afford, throughput bottleneck và purchase cadence; giá để [TBD].", "Giảm thời gian chờ và giữ động lực ở giữa hành trình.", "PRG-01; ECO-01; ANA-02", "P0", "L", "Trong >= [số người] playtest, P50 thời gian giữa hai meaningful purchase ở mid-game nằm trong [khoảng phút]; tỷ lệ session rời tại cùng một node < [ngưỡng %]; không có đoạn chờ > [phút] khi chơi đúng loop."),
    item("PRG-05", "1. Progression Expansion", "Late-game pacing", "Tune các node cuối, landmark và bước chuẩn bị prestige; không đặt giá/reward cụ thể trong backlog.", "Tạo cao trào rõ ràng trước khi hoàn thành nhà hàng.", "PRG-04; ECO-05; ANA-03", "P0", "L", "Trong >= [số người] playtest, P50 late-game nằm trong [khoảng phút]; >= [tỷ lệ %] người test hiểu mục tiêu cuối mà không cần admin giải thích; không có purchase gap > [phút]."),
    item("PRG-06", "1. Progression Expansion", "Điều kiện hoàn thành nhà hàng", "Định nghĩa server-authoritative completion rule từ unlock bắt buộc, build revision và trạng thái operation.", "Cho người chơi đích đến minh bạch và đáng tin cậy.", "PRG-01; PRG-03", "P0", "M", "Completion chỉ bật khi 100% điều kiện bắt buộc đã commit và apply; thiếu bất kỳ điều kiện nào đều bị từ chối với lý do; 10/10 test thứ tự mua khác nhau cho cùng kết quả."),
    item("PRG-07", "1. Progression Expansion", "Mục tiêu prestige lần đầu", "Đo và tune first completion/first prestige ở 150-210 phút active play, phân đoạn solo/co-op.", "Đảm bảo hành trình MVP đủ dài nhưng vẫn có thể hoàn thành.", "PRG-06; PRE-01; ANA-01", "P0", "L", "Với >= [cỡ mẫu] người mới, P50 solo active time từ claim đến đủ điều kiện prestige nằm trong 150-210 phút; co-op không thấp hơn 120 phút; dashboard tách solo/2/3/4 người."),

    # 2. Restaurant Content
    item("CON-01", "2. Restaurant Content", "Restaurant area", "Chốt danh sách khu vực [TBD] và anchor bắt buộc cho entrance, queue, table, station, register, exit.", "Nhà hàng mở rộng trực quan và dễ định hướng.", "PRG-03", "P0", "L", "100% area có bounds và anchor hợp lệ; route validator đi được entrance-to-exit; không anchor nào ngoài plot; 10/10 lần load giữ đúng vị trí."),
    item("CON-02", "2. Restaurant Content", "Recipe", "Cấu hình danh sách recipe [TBD], version, station sequence và điều kiện unlock; không ghi recipe cụ thể.", "Mở rộng lựa chọn order và hoạt động bếp.", "CON-03; PRG-01", "P1", "L", "Có [số lượng] recipe được duyệt; 100% recipe hoàn thành được bằng station đã unlock; version cũ của order vẫn settle đúng sau reload; validator chạy 0 lỗi."),
    item("CON-03", "2. Restaurant Content", "Cooking station", "Cấu hình station [TBD], capacity, queue, interaction point và job timeout.", "Tạo bottleneck và nâng cấp bếp dễ hiểu.", "CON-01; PER-04", "P0", "L", "100% station có ID/capacity/anchor; station full từ chối job mới; 20 concurrent claim test chỉ có 1 winner/job slot; restart không mất job committed."),
    item("CON-04", "2. Restaurant Content", "Customer type", "Định nghĩa customer type [TBD] bằng preference, patience profile và điều kiện xuất hiện; không tự chọn chủ đề.", "Tạo biến thể nhu cầu và nhịp phục vụ.", "CUS-03; CON-02", "P1", "M", "Có [số lượng] type hợp lệ; mỗi type chỉ chọn recipe được phép; phân phối spawn trong 1.000 mẫu lệch <= [ngưỡng %] so với cấu hình; unknown type fallback an toàn."),
    item("CON-05", "2. Restaurant Content", "Staff role", "Cấu hình role Host/Waiter/Prep Cook/Grill Cook/Cashier/Cleaner với responsibility và unlock mapping.", "Cho thấy nhà hàng dần vận hành chuyên nghiệp hơn.", "STF-01..STF-06; PRG-01", "P0", "M", "Đủ 6 role có definition; mỗi responsibility có đúng 1 owner logic tại một thời điểm; role bị khóa không chạy automation; validator role mapping báo 0 lỗi."),
    item("CON-06", "2. Restaurant Content", "Upgrade", "Định nghĩa upgrade [TBD] cho capacity, speed hoặc quality bằng bounded modifier và version.", "Cung cấp lựa chọn cải thiện hiệu suất có tác động rõ.", "ECO-04; PRG-01", "P1", "L", "100% upgrade có before/after metric; modifier stacking không vượt [cap]; mua lại cùng upgrade không nhân hiệu ứng; 10/10 reload giữ đúng level."),
    item("CON-07", "2. Restaurant Content", "Cosmetic", "Lập danh mục cosmetic [TBD] tách khỏi power progression và economy competitive.", "Tăng cá tính và mục tiêu sưu tầm không pay-to-win.", "PRG-01", "P2", "M", "100% cosmetic không đổi revenue/throughput/patience; ownership persist qua prestige/restart; 10/10 permission test ngăn dùng cosmetic chưa sở hữu."),
    item("CON-08", "2. Restaurant Content", "Visual milestone", "Định nghĩa các mốc thay đổi silhouette/geometry [TBD] gắn với major purchase.", "Người chơi và visitor nhìn thấy tiến bộ ngay lập tức.", "CON-01; PRG-03", "P1", "L", "Mỗi progression band có >= 1 visual milestone; >= [tỷ lệ %] tester nhận ra before/after trong blind test; mỗi milestone apply trong <= [giây] theo build budget."),
    item("CON-09", "2. Restaurant Content", "Sound và animation", "Tạo cue âm thanh/animation ngắn cho purchase, cooking, service, error và prestige; asset cụ thể [TBD].", "Làm hành động rõ phản hồi và tăng cảm giác hoàn thiện.", "CON-03; PRG-03; PRE-02", "P1", "L", "100% critical action có success/error cue phân biệt; animation purchase dài 1-4 giây; không vượt [số] display/particle event mỗi plot/giây; 0 warning asset thiếu khi load."),

    # 3. Customer Simulation
    item("CUS-01", "3. Customer Simulation", "Queue", "Triển khai logical queue có capacity, fixed anchor và backpressure khi hết table/entity budget.", "Khách đến có trật tự và bottleneck dễ quan sát.", "CON-01; PERF-02", "P0", "L", "Queue không vượt configured capacity trong 1.000 arrival; full queue không spawn entity mới; thứ tự FIFO đúng >= 99% trừ priority rule được cấu hình; deactivation xóa 100% entity queue."),
    item("CUS-02", "3. Customer Simulation", "Table reservation", "Dùng reservation ID và compare-and-set để một table chỉ thuộc một party.", "Ngăn khách trùng bàn và bàn bị khóa vĩnh viễn.", "CUS-01; PER-04", "P0", "L", "Trong 100 race test, mỗi table có tối đa 1 active reservation; reservation timeout giải phóng trong <= [giây]; restart reconciliation không để orphan reservation."),
    item("CUS-03", "3. Customer Simulation", "Customer state machine", "Hoàn thiện state machine SPAWNING đến LEAVING/ABANDONED/TECHNICAL_FAILURE với revision guard.", "Luồng phục vụ nhất quán, dễ hiểu và có thể phục hồi.", "CUS-01; CUS-02; PER-04", "P0", "XL", "100% transition không hợp lệ bị từ chối; toàn bộ state có entry/exit test; 10.000 simulated order không có stuck state quá timeout; state revision luôn tăng đơn điệu."),
    item("CUS-04", "3. Customer Simulation", "Patience", "Tách patience cho queue, order và food wait; cấu hình giá trị [TBD] và feedback theo bottleneck.", "Cho phản hồi công bằng về điểm nghẽn vận hành.", "CUS-03; CON-04", "P1", "M", "Đủ 3 patience timer độc lập; mỗi timer pause/resume đúng theo rule; 50 test boundary tại timeout cho kết quả xác định; technical failure không trừ satisfaction."),
    item("CUS-05", "3. Customer Simulation", "Timeout", "Định nghĩa timeout và recovery action cho mọi customer/order/table state.", "Không để nhà hàng bị kẹt do lỗi hoặc người chơi rời đi.", "CUS-03; CUS-04", "P0", "L", "100% state có timeout hoặc lý do exempt được duyệt; injected timeout ở từng state phục hồi trong <= [giây]; không tạo payout ngoài SETTLED; 0 table kẹt sau suite."),
    item("CUS-06", "3. Customer Simulation", "Pathfinding recovery", "Thêm no-progress watchdog, retry budget, safe teleport/despawn và release reservation.", "Lỗi đường đi không làm mất doanh thu hợp lệ hay khóa bàn.", "CUS-03; PERF-05", "P0", "L", "100/100 injected no-path được phát hiện trong <= [giây]; recovery không giảm rating; reservation được release/reconcile; path retry không vượt [số]/customer/phút."),
    item("CUS-07", "3. Customer Simulation", "Plot observed/unobserved behavior", "Tách observed physical simulation, online-unobserved batch và offline zero-income cho MVP.", "Duy trì tiến trình online hợp lý mà không gây tải NPC dư thừa.", "CUS-03; ECO-07; PERF-03", "P0", "XL", "Unobserved/offline có 0 customer entity sau <= [giây]; online-unobserved chỉ tạo base automation income; offline tạo 0 income; 100 transition qua lại không duplicate order/payout."),
    item("CUS-08", "3. Customer Simulation", "Entity cleanup", "Quản lý ownership tag và cleanup khi deactivate, unload, reset, reassignment, disable và restart.", "Ngăn lag, NPC ma và trạng thái sai.", "CUS-07; VIS-06; ADM-06", "P0", "L", "Sau mỗi lifecycle test, plugin-owned customer entity còn lại = 0 trên inactive plot; không leak task/handle/chunk ticket; stale entity load bị xóa trong <= [tick]."),

    # 4. Staff and Automation
    item("STF-01", "4. Staff và Automation", "Host", "Tự động gán khách chờ vào table phù hợp bằng reservation CAS.", "Giảm thao tác lặp ở cửa vào khi nhà hàng đông.", "CUS-01; CUS-02; CON-05", "P1", "M", "Host chỉ gán table AVAILABLE phù hợp; 100 race test không double-reserve; khi disabled không tự gán; assignment latency p95 <= [giây]."),
    item("STF-02", "4. Staff và Automation", "Waiter", "Tự động nhận order và [tùy chọn: giao món] theo level đã chốt.", "Giảm thao tác phục vụ lặp lại nhưng giữ loop chủ động.", "CUS-03; CON-05", "P1", "L", "Waiter chỉ xử lý order hợp lệ trong restaurant; một order có tối đa 1 assignment; 100 concurrent test không claim trùng; behavior theo level đúng 100% test case."),
    item("STF-03", "4. Staff và Automation", "Prep Cook", "Tự động xử lý các bước prep/assembly được cấu hình với bounded queue.", "Tăng throughput và làm rõ chuyên môn hóa bếp.", "CON-02; CON-03; CON-05", "P1", "L", "Prep Cook chỉ chạy step đúng loại; không vượt station capacity; mỗi job completion đúng một lần; 10/10 restart giữa job giữ kết quả chính xác."),
    item("STF-04", "4. Staff và Automation", "Grill Cook", "Tự động xử lý station cook bằng absolute finishes_at và claim guard.", "Giải phóng thao tác nấu lặp lại ở mid/late game.", "CON-02; CON-03; CON-05", "P0", "L", "Grill Cook không tạo job thiếu input/order; finish timestamp sai lệch <= [giây]; output entitlement tối đa 1/job trong 100 retry/race test."),
    item("STF-05", "4. Staff và Automation", "Cashier", "Tự động collect/auto-deposit theo level, nhưng settlement vẫn qua ledger transaction.", "Giảm việc chạy về quầy thu tiền và làm progression mượt hơn.", "ECO-02; ECO-09; CON-05", "P0", "L", "Mỗi SETTLED order tạo đúng 1 ledger credit; Cashier disabled không auto-collect; 100 disconnect/retry test không mất hoặc nhân đôi cash."),
    item("STF-06", "4. Staff và Automation", "Cleaner", "Tự động chuyển table DIRTY qua CLEANING về AVAILABLE theo capacity.", "Giảm bottleneck vệ sinh ở nhà hàng lớn.", "CUS-02; CUS-03; CON-05", "P1", "M", "Cleaner chỉ claim table DIRTY; mỗi table có tối đa 1 cleaning job; 100 race test không double-clean; timeout/restart trả table về state hợp lệ."),
    item("STF-07", "4. Staff và Automation", "Staff upgrade", "Cấu hình level, effect, cap và pricing placeholder cho từng staff role.", "Cho người chơi tiếp tục tối ưu sau khi đã thuê staff.", "STF-01..STF-06; ECO-04", "P1", "L", "Mỗi level có effect đo được và bounded; không mua vượt max level; 10/10 concurrent purchase chỉ tăng 1 level/lần; level persist qua restart."),
    item("STF-08", "4. Staff và Automation", "Automation balance", "Cân bằng automated throughput theo progression và quy mô co-op; không cho worker tăng spawn cap.", "Automation có giá trị nhưng không phá pacing hoặc economy.", "STF-01..STF-07; ECO-07; ANA-06", "P0", "L", "Ở cùng progression, automated baseline nằm trong [dải mục tiêu]; 4-player throughput <= 1,25x solo cap; 60 phút simulation không vượt currency budget [TBD]."),
    item("STF-09", "4. Staff và Automation", "Active gameplay sau automation", "Cung cấp intervention 30-90 giây theo cadence 4-7 phút, dùng capped reward budget.", "Giữ người chơi có quyết định ý nghĩa sau khi tự động hóa.", "STF-08; ECO-06; ANA-05", "P0", "L", "Observed active plot nhận 1 optional intervention mỗi 4-7 phút theo cấu hình; bỏ qua không dừng base income; active uplift 15-25% mục tiêu và không vượt hard cap 30%/15 phút."),

    # 5. Co-op
    item("COP-01", "5. Co-op", "Mời thành viên", "Cho owner mời, accept, revoke và giới hạn tổng nhóm 2-4 người.", "Bạn bè có thể cùng vận hành một nhà hàng.", "PER-03; COP-02", "P0", "L", "Nhà hàng có tối đa 1 owner + 3 collaborator; invite hết hạn sau [thời gian]; duplicate accept không tạo member trùng; 50 race invite/accept giữ đúng giới hạn."),
    item("COP-02", "5. Co-op", "Role và permission", "Định nghĩa role [TBD] và permission riêng cho spend, staff, build, invite; prestige owner-only.", "Hợp tác an toàn với quyền kiểm soát rõ ràng.", "PER-01", "P0", "XL", "100% protected action kiểm tra role + permission revision server-side; collaborator không thể prestige trong 50 bypass test; thay quyền có hiệu lực trong <= [giây]."),
    item("COP-03", "5. Co-op", "Chia công việc", "Dùng assignment CAS cho order, station, table và cleaning để nhiều người không claim cùng việc.", "Cho nhóm phối hợp hiệu quả mà không xung đột.", "CUS-03; CON-03; COP-02", "P0", "XL", "Trong 100 race test/action type chỉ có 1 winner; loser nhận feedback trong <= [ms]; disconnect release/reassign theo rule; không mất committed progress."),
    item("COP-04", "5. Co-op", "Reward budget", "Mỗi order/event có fixed reward budget; phân chia participation không tạo thêm currency.", "Co-op hữu ích nhưng không bắt buộc để tối ưu economy.", "ECO-06; COP-03", "P0", "L", "Tổng reward của 1 order qua 1-4 người <= fixed budget trong 1.000 simulation; rounding không tạo đơn vị dư; 4-player throughput kinh tế <= 1,25x solo cap."),
    item("COP-05", "5. Co-op", "Chống nhân đôi reward", "Ràng buộc payout bằng order/event ID, participant và idempotency key.", "Bảo vệ economy khỏi spam/reconnect/race exploit.", "COP-04; ECO-09", "P0", "XL", "100.000 retry/concurrent payout tạo tối đa 1 ledger result/logical reward; unique constraint bắt duplicate; balance reconciliation chênh lệch = 0."),
    item("COP-06", "5. Co-op", "Thành viên disconnect", "Quy định release assignment, dish entitlement và session token khi collaborator rời.", "Nhóm tiếp tục chơi mà không bị khóa công việc.", "COP-03; PER-04", "P0", "L", "Disconnect ở từng active state được recover trong <= [giây]; không duplicate dish/reward; assignment không còn session được release hoặc chuyển về reconciliation."),
    item("COP-07", "5. Co-op", "Owner offline", "Dừng nhận khách mới, drain/settle việc committed rồi quiesce; visitor/collaborator không bật income vô hạn.", "Trạng thái nhà hàng nhất quán khi chủ rời server.", "CUS-07; COP-06; ECO-08", "P0", "XL", "Sau owner offline, new arrival = 0 trong <= [giây]; committed SETTLED không mất; plot quiesce trong <= [phút]; offline income và active bonus đều = 0."),
    item("COP-08", "5. Co-op", "Co-op audit log", "Log invite, role, permission, spend, refund, assignment quan trọng và actor/session/correlation ID.", "Giải quyết tranh chấp và hỗ trợ điều tra exploit.", "COP-02; ANA-08", "P1", "M", "100% mutation co-op thuộc danh sách audit có actor, target, before/after, permission revision, timestamp, correlation ID; truy vấn 30 ngày trả kết quả trong <= [giây]."),

    # 6. Visit System
    item("VIS-01", "6. Visit System", "Thăm nhà hàng", "Thêm `/restaurant visit <player>` và [social browser TBD] với kiểm tra plot availability.", "Khuyến khích khám phá và khoe tiến trình nhà hàng.", "PER-03; CON-01", "P1", "L", "Visit thành công đưa người chơi vào đúng safe anchor trong <= [giây]; invalid/offline target trả lỗi rõ; 50 simultaneous visit không đưa người chơi ra ngoài bounds."),
    item("VIS-02", "6. Visit System", "Visitor permission", "Định nghĩa visitor role read-only và allowlist interaction xã hội [TBD].", "Cho phép tham quan an toàn mà không ảnh hưởng chủ nhà.", "VIS-01; COP-02", "P0", "M", "100% visitor nhận role read-only khi enter; role bị thu hồi khi leave; 50 permission bypass test không mutate cash/unlock/order/staff."),
    item("VIS-03", "6. Visit System", "Chống visitor sử dụng purchase control", "Kiểm tra membership, permission revision và plot/control ID tại mọi purchase boundary.", "Bảo vệ tiền và progression của chủ nhà.", "VIS-02; ECO-09", "P0", "L", "Visitor bị từ chối 100/100 click/packet/replay test; không ledger row, unlock hoặc world operation được tạo; attempt có structured security event."),
    item("VIS-04", "6. Visit System", "Hiển thị tiến trình", "Hiển thị completion, prestige và visual milestone đã công khai; dữ liệu nhạy cảm [TBD] không lộ.", "Visitor hiểu thành tựu của nhà hàng đang xem.", "VIS-01; PRG-06", "P2", "M", "UI phản ánh đúng committed progression revision trong <= [giây]; 10/10 privacy test không hiển thị balance/audit/private setting; stale cache tự refresh."),
    item("VIS-05", "6. Visit System", "Giới hạn tương tác", "Chặn break/place, container, station, entity damage, vehicle, item pickup/drop và boundary exploit cho visitor.", "Tham quan không thể phá hoặc làm nghẽn nhà hàng.", "VIS-02; CON-01", "P0", "L", "Toàn bộ interaction matrix [TBD] có automated test; 100% destructive/economic action bị chặn; action xã hội allowlist hoạt động; visitor không giữ item entitlement khi leave."),
    item("VIS-06", "6. Visit System", "Visitor cleanup", "Dọn session, temporary UI/entity, task và teleport visitor ra ngoài khi plot reset/reassign/close.", "Không để visitor mắc kẹt hoặc tạo tài nguyên rò rỉ.", "VIS-01; CUS-08", "P0", "M", "Sau leave/reset/disable, visitor session và temporary entity/task còn lại = 0; evacuation hoàn tất trong <= [giây]; reconnect không khôi phục quyền visitor cũ."),

    # 7. Prestige
    item("PRE-01", "7. Prestige", "Điều kiện prestige", "Kiểm tra completion, owner identity, no active blocking operation và prestige version.", "Đảm bảo prestige là thành tựu cuối vòng chơi hợp lệ.", "PRG-06; PER-05", "P0", "L", "Prestige chỉ available khi 100% condition đúng; mỗi condition sai trả reason code; collaborator/visitor bị từ chối; 20 ordering test không bypass gate."),
    item("PRE-02", "7. Prestige", "Prestige confirmation", "Dùng confirmation hai bước, hiển thị reset/retain summary và operation token hết hạn.", "Ngăn reset nhầm và giúp quyết định có thông tin.", "PRE-01", "P0", "M", "Không prestige bằng một click; confirmation hiển thị 100% category reset/retain; token hết hạn sau [thời gian]; click lặp chỉ khởi tạo 1 operation."),
    item("PRE-03", "7. Prestige", "Reset progression", "Reset cash, local unlock/build/staff theo durable fenced state machine.", "Mở vòng tăng trưởng mới nhất quán và an toàn.", "PRE-02; PER-05", "P0", "XL", "Sau completion, 100% field reset khớp rule; world trở về canonical base revision; retry/restart tại từng phase hội tụ cùng trạng thái; không còn entity/task cũ."),
    item("PRE-04", "7. Prestige", "Giữ lại reward phù hợp", "Định nghĩa retain matrix cho cosmetic, achievement, social identity và [reward TBD].", "Người chơi giữ cảm giác thành tựu qua mỗi vòng.", "PRE-03; CON-07", "P0", "L", "100% retained field giữ nguyên qua 10 prestige/restart test; non-retained field về default; matrix version được snapshot trong prestige history."),
    item("PRE-05", "7. Prestige", "Prestige bonus", "Áp dụng bounded bonus và unlock route [TBD], tránh exponential stacking.", "Vòng tiếp theo nhanh hơn vừa phải và có điểm mới.", "PRE-04; ECO-05", "P1", "L", "Bonus không vượt configured cap ở [số] prestige level; revenue calculation dùng đúng snapshot; early-stage time không thấp hơn [ngưỡng] trong balance simulation."),
    item("PRE-06", "7. Prestige", "Anti-duplicate", "Dùng prestige operation ID, expected level/revision và unique constraint.", "Ngăn nhân đôi bonus, level hoặc reset exploit.", "PRE-03; ECO-09", "P0", "XL", "100.000 concurrent/retry request chỉ tăng prestige level đúng 1; chỉ có 1 history row/operation; duplicate bị từ chối mà không đổi balance/unlock."),
    item("PRE-07", "7. Prestige", "Recovery khi restart", "Resume/reconcile prestige ở mọi phase sau server crash hoặc plugin restart.", "Không mất nhà hàng hoặc bị kẹt giữa hai vòng chơi.", "PRE-03; PER-06", "P0", "XL", "Crash injection tại 100% phase được recover tự động; trạng thái cuối khớp expected reset/retain; recovery hoàn tất trong <= [phút]; không cần sửa DB thủ công."),
    item("PRE-08", "7. Prestige", "Audit log", "Ghi actor, old/new level, reset/retain version, operation state và correlation ID.", "Cho phép support xác minh và khôi phục sự cố prestige.", "PRE-03; ADM-05", "P0", "M", "100% attempt success/fail có audit event; event chứa đủ [danh sách field]; query theo restaurant/operation trong <= [giây]; log không chứa secret."),

    # 8. Economy
    item("ECO-01", "8. Economy", "Balance sheet", "Lập bảng theo progression band: CPM active/auto, order rate, price placeholder, time-to-afford, payback và solo/co-op.", "Tạo pacing nhất quán và tránh đoạn chờ vô nghĩa.", "PRG-01", "P0", "XL", "100% purchase node có price input và time-to-afford projection; đủ scenario solo/2/4 người; tổng projected P50 first prestige nằm 150-210 phút; sheet không có ô tính bắt buộc bị trống."),
    item("ECO-02", "8. Economy", "Currency source", "Đăng ký mọi nguồn currency bằng operation type, cap, related ID và eligibility.", "Thu nhập dễ hiểu và được kiểm soát.", "ECO-01; PER-03", "P0", "L", "100% credit ledger thuộc source allowlist; unknown source bị từ chối; source report khớp tổng ledger với sai số 0; không source nào dùng floating point."),
    item("ECO-03", "8. Economy", "Currency sink", "Đăng ký purchase, staff, upgrade, prestige và cosmetic sink; nội dung/giá cụ thể [TBD].", "Tiền luôn có mục đích sử dụng và progression có lựa chọn.", "ECO-01; ECO-09", "P0", "L", "100% debit thuộc sink allowlist; balance không âm; insufficient funds không tạo unlock/world operation; sink report khớp ledger với sai số 0."),
    item("ECO-04", "8. Economy", "Upgrade pricing", "Tune price curve và expected payback theo band, để giá cụ thể trong input [TBD].", "Mỗi nâng cấp có chi phí hợp lý so với tác động.", "ECO-01; CON-06", "P0", "L", "100% functional upgrade có payback projection; median payback nằm [khoảng phút] đã duyệt; không node nào affordable trước prerequisite; pricing version được snapshot lúc mua."),
    item("ECO-05", "8. Economy", "Revenue scaling", "Dùng base value, bounded station/prestige modifier và capped additive bonus.", "Tăng trưởng rõ nhưng không làm các giai đoạn đầu mất ý nghĩa.", "ECO-01; PRE-05", "P0", "L", "10.000 calculation test khớp công thức; mọi multiplier <= cap; dùng BIGINT và rounding rule xác định; không overflow ở max configured input."),
    item("ECO-06", "8. Economy", "Manual bonus", "Đặt active service/event bonus trong fixed window budget; giá trị cụ thể [TBD].", "Khuyến khích chơi chủ động mà không ép grind.", "ECO-05; STF-09", "P0", "M", "Active uplift mục tiêu 15-25%/15 phút và hard cap <= 30%; spam 10x action không vượt cap; unobserved/AFK không nhận bonus; report theo reason code."),
    item("ECO-07", "8. Economy", "Automation income", "Tính base automated throughput theo unlock/staff/capacity và simulation state.", "Nhà hàng vẫn tiến triển online khi người chơi quan sát ít hơn.", "STF-08; CUS-07", "P0", "L", "Cùng snapshot cho kết quả deterministic; online-unobserved chỉ có base income; physical và batch không cùng advance một order; 1.000 transition không duplicate payout."),
    item("ECO-08", "8. Economy", "AFK rule", "Định nghĩa activity signal server-side và loại active/VIP/co-op/leaderboard reward khi AFK.", "Giữ economy công bằng và giảm macro exploit.", "ECO-06; ECO-07", "P0", "L", "AFK sau [thời gian] không nhận restricted reward; movement loop/held control không reset activity trong test; base automation tuân đúng state; 100% transition có analytics event."),
    item("ECO-09", "8. Economy", "Transaction idempotency", "Mọi payout/purchase/refund/grant dùng stable idempotency key, atomic ledger và revision.", "Bảo vệ tiền chính xác qua spam, lag và restart.", "PER-01; PER-02", "P0", "XL", "100.000 duplicate/retry test cho mỗi operation type tạo tối đa 1 mutation; balance + ledger + related state commit atomically; reconciliation chênh lệch = 0."),
    item("ECO-10", "8. Economy", "Economy analytics", "Xuất source/sink, balance delta, progression band, co-op size, active/auto và anomaly event.", "Cho team phát hiện inflation, bottleneck và exploit sớm.", "ECO-02..ECO-09; ANA-08", "P1", "M", "Event coverage >= 99,9% so với ledger trong test; dashboard lọc theo band/co-op/state; daily source-sink delta khớp ledger; anomaly threshold [TBD] tạo alert."),

    # 9. Persistence and Recovery
    item("PER-01", "9. Persistence và Recovery", "PostgreSQL", "Dùng PostgreSQL + HikariCP làm durable source of truth cho production validation.", "Tiến trình và tiền an toàn, hỗ trợ vận hành thực tế.", "[Chọn PostgreSQL version]; [Chọn hosting]", "P0", "XL", "Server khởi động và chạy test suite với PostgreSQL version đã pin; synchronous SQL trên main thread = 0; pool metric/timeout được cấu hình; DB unavailable vào degraded mode."),
    item("PER-02", "9. Persistence và Recovery", "Schema migration", "Version hóa schema bằng [Flyway/runner TBD], forward migration và compatibility check.", "Update server không làm mất hoặc làm sai dữ liệu.", "PER-01", "P0", "L", "Fresh install và upgrade từ Phase 1 đều pass; chạy migration lại không đổi dữ liệu; failure dừng startup an toàn; checksum mismatch có lỗi rõ."),
    item("PER-03", "9. Persistence và Recovery", "Cash/unlock persistence", "Lưu balance BIGINT, ledger, unlock, price/definition snapshot và revision atomically.", "Người chơi không mất tiền hoặc nâng cấp sau reconnect/restart.", "PER-01; PER-02; ECO-09", "P0", "XL", "1.000 purchase/restart test giữ balance/unlock chính xác; balance không âm; mỗi unlock unique/restaurant; ledger sum và current balance chênh lệch = 0."),
    item("PER-04", "9. Persistence và Recovery", "Order persistence", "Lưu business state cho order, reservation, station job, entitlement và settlement; không lưu từng vị trí NPC.", "Order đang chạy có thể phục hồi mà không nhân đôi món hoặc payout.", "PER-01; CUS-03; CON-03", "P0", "XL", "Crash tại mỗi order phase phục hồi theo policy; một order tối đa 1 payout và 1 active entitlement; 100% orphan record được reconcile; không cần persist transient coordinates."),
    item("PER-05", "9. Persistence và Recovery", "Build projection", "Commit purchase/world_operation trong DB rồi apply canonical stage với plot fence token.", "Đã trả tiền luôn hội tụ về đúng công trình, kể cả sau crash.", "PER-03; PRG-03", "P0", "XL", "Crash sau commit/trước apply tạo repair và hoàn tất sau restart; retry 3 lần cho world hash giống nhau; stale fence token không sửa world; operation audit đầy đủ."),
    item("PER-06", "9. Persistence và Recovery", "Restart recovery", "Khôi phục restaurant run, job, entitlement, world operation, reset và prestige theo durable state.", "Restart có kế hoạch không làm mất hoặc nhân đôi tiến trình.", "PER-03..PER-05; PRE-07", "P0", "XL", "Restart tại 100% checkpoint suite hội tụ trong <= [phút]; balance/unlock/order/world revision đúng; leaked entity/task/chunk ticket = 0; player nhận thông báo recovery."),
    item("PER-07", "9. Persistence và Recovery", "Crash recovery", "Chạy crash-injection sau các commit boundary và xử lý in-flight command bằng stable operation ID.", "Sự cố đột ngột không phá economy hoặc plot.", "PER-06; TST-05", "P0", "XL", "Kill process tại >= [số checkpoint] cho purchase/order/prestige/reset; sau boot chênh lệch ledger = 0; không duplicate unlock/payout/bonus; mọi operation đạt terminal/retryable state."),
    item("PER-08", "9. Persistence và Recovery", "Reconciliation", "Đối chiếu balance-ledger, unlock-stage, job-entitlement, assignment-fence và entity ownership.", "Tự phát hiện và sửa sai lệch trước khi ảnh hưởng nhiều người chơi.", "PER-03..PER-07; ADM-07", "P0", "XL", "Phát hiện 100% fixture sai lệch đã định nghĩa; auto-fix chỉ áp dụng rule an toàn; unresolved issue tạo alert/audit; rerun sau fix cho 0 discrepancy."),
    item("PER-09", "9. Persistence và Recovery", "Backup và restore test", "Chốt RPO/RTO [TBD], backup DB/world/artifact/content và diễn tập restore + reconciliation.", "Bảo vệ công sức người chơi trước mất dữ liệu nghiêm trọng.", "PER-01; PER-08; [Backup storage TBD]", "P0", "XL", "Restore drill trên staging đạt RPO <= [mục tiêu] và RTO <= [mục tiêu]; checksum artifact đúng; reconciliation sau restore chênh lệch cash/unlock = 0; biên bản drill được lưu."),

    # 10. Analytics
    item("ANA-01", "10. Analytics", "Session duration", "Track session start/end, active time, AFK time, restaurant ID và co-op size.", "Cho team tune mục tiêu 150-210 phút bằng dữ liệu thật.", "ECO-08; [Analytics sink TBD]", "P0", "M", ">= 99% session test có start/end hoặc timeout close; active time không tính AFK; duration không âm; dashboard hiển thị P50/P80 theo solo/2/3/4 người."),
    item("ANA-02", "10. Analytics", "Purchase funnel", "Track node viewed/affordable/purchased và time since previous purchase theo definition version.", "Xác định nơi người chơi chậm lại hoặc bỏ cuộc.", "PRG-01; ECO-09", "P0", "M", "100% successful purchase có event và node/version; funnel từ first node đến prestige truy vấn được; event duplicate rate < [ngưỡng %]; đối chiếu purchase table coverage >= 99,9%."),
    item("ANA-03", "10. Analytics", "Progression bottleneck", "Tính dwell time, insufficient-funds attempt, queue/station saturation và drop-off theo band.", "Giúp sửa đúng điểm gây chờ hoặc khó hiểu.", "ANA-01; ANA-02; CUS-04", "P1", "M", "Dashboard xếp hạng top [N] bottleneck theo dwell/drop-off; lọc theo branch và co-op; dữ liệu cập nhật trong <= [độ trễ]; fixture known bottleneck đứng đúng top."),
    item("ANA-04", "10. Analytics", "Recipe usage", "Track recipe selected, completed, abandoned, technical failure và version; không lưu nội dung chat.", "Cho biết content nào được dùng và nơi flow bếp thất bại.", "CON-02; PER-04", "P1", "S", ">= 99% order terminal có recipe outcome event; tổng outcome khớp order table trong sai số [ngưỡng]; dashboard lọc version/customer type; không có PII ngoài UUID cần thiết."),
    item("ANA-05", "10. Analytics", "Active job usage", "Track manual take/cook/serve/clean/intervention, duration và capped bonus.", "Đo việc gameplay chủ động còn hấp dẫn sau automation.", "STF-09; ECO-06", "P1", "M", "Mỗi active job success/fail có reason code; tỷ lệ participation sau automation tính được; bonus event khớp ledger; spam rejected không được tính là active completion."),
    item("ANA-06", "10. Analytics", "Automation usage", "Track staff enabled/level, automated jobs, throughput và observed/unobserved state.", "Giúp cân bằng thời điểm automation và hiệu quả từng role.", "STF-07; ECO-07", "P1", "M", "Automated completion được gắn đúng staff role/level/state; dashboard so sánh manual/auto; event count khớp job records >= 99%; không ghi event mỗi tick."),
    item("ANA-07", "10. Analytics", "Prestige conversion", "Track eligibility, confirmation open/cancel/complete, active minutes và branch.", "Đo tỷ lệ người chơi đạt và thực hiện vòng lặp cốt lõi.", "PRE-01; PRE-08; ANA-01", "P0", "M", "100% prestige attempt có funnel event; time-to-eligible và time-to-complete tính được; dashboard hiển thị conversion/P50/P80 theo branch và co-op; duplicate event được dedupe."),
    item("ANA-08", "10. Analytics", "Error/recovery event", "Chuẩn hóa event cho timeout, path recovery, reconciliation, degraded mode, duplicate rejection và admin repair.", "Cho team phát hiện vấn đề ổn định và exploit nhanh.", "PER-08; CUS-06; ECO-09", "P0", "L", "100% critical error path có code, severity, restaurant, operation/correlation ID; không log secret; alert tạo trong <= [phút] khi vượt threshold [TBD]; test event schema pass."),

    # 11. Admin Tools
    item("ADM-01", "11. Admin Tools", "Player/plot inspect", "Lệnh/UI read-only hiển thị owner/member, run, balance revision, unlock, order, operation và entity count.", "Support chẩn đoán nhanh mà không sửa nhầm dữ liệu.", "PER-03..PER-08", "P0", "L", "Inspect trả dữ liệu trong <= [giây] cho plot test; khớp DB/runtime revision; không lộ secret; user thiếu permission bị từ chối và audit."),
    item("ADM-02", "11. Admin Tools", "Reset state machine", "Cho admin khởi động durable fenced reset: QUIESCING đến RECONCILING/COMPLETE.", "Khôi phục plot hỏng bằng quy trình an toàn.", "PER-05; COP-07", "P0", "XL", "Reset không chạy song song purchase/build; restart ở từng phase resume an toàn; stale fence bị từ chối; complete để lại canonical stage và 0 entity/task cũ."),
    item("ADM-03", "11. Admin Tools", "Repair plot", "Replay committed stage và rebuild controls/displays mà không tự thay đổi unlock/cash.", "Sửa hình ảnh nhà hàng mà không ảnh hưởng progression.", "PER-05; PER-08", "P0", "L", "Repair đưa world hash về canonical target; balance/unlock trước và sau không đổi; chạy lại 3 lần idempotent; validation report lưu vào audit."),
    item("ADM-04", "11. Admin Tools", "Grant/refund", "Thao tác có reason bắt buộc, idempotency key, approval rule [TBD] và ledger immutable.", "Support xử lý sự cố công bằng và có thể truy vết.", "ECO-09; ADM-08", "P0", "L", "100% grant/refund có amount, reason, actor, target, related ID và correlation ID; retry không trả hai lần; permission test chặn unauthorized; balance-ledger chênh lệch = 0."),
    item("ADM-05", "11. Admin Tools", "Economy audit", "Tra cứu ledger theo restaurant/player/time/reason/operation và chạy invariant report.", "Phát hiện exploit và giải quyết khiếu nại tiền tệ.", "ECO-10; PER-08", "P0", "M", "Query [số lượng] ledger row hoàn tất trong <= [giây]; report xác nhận revision liên tục và sum đúng; export redacted; discrepancy tạo case ID."),
    item("ADM-06", "11. Admin Tools", "Entity cleanup", "Quét/removal plugin-owned stale customer/display/interaction theo plot và ownership tag.", "Giảm lag và loại bỏ NPC/control ma.", "CUS-08; VIS-06", "P0", "M", "Dry-run liệt kê đúng fixture; execute xóa 100% stale và 0 valid entity; cleanup hoàn tất trong budget [ms/tick]; mọi removal có audit summary."),
    item("ADM-07", "11. Admin Tools", "Force reconciliation", "Cho phép chạy reconciliation scoped theo restaurant/plot/operation với dry-run mặc định.", "Khôi phục dữ liệu nhất quán mà không thao tác DB thủ công.", "PER-08; ADM-08", "P0", "L", "Dry-run không mutate; execute chỉ áp dụng safe rule; rerun cho 0 safe discrepancy; unresolved issue giữ nguyên và tạo audit/alert; cần permission riêng."),
    item("ADM-08", "11. Admin Tools", "Permission và audit log", "Phân quyền riêng cho inspect/reset/repair/grant/refund/cleanup/reconcile và log mọi attempt.", "Ngăn lạm dụng công cụ có quyền cao.", "[Permission provider TBD]; ANA-08", "P0", "L", "100% admin action kiểm tra permission server-side; success/failure đều audit actor, target, reason, before/after, correlation ID; log append-only; retention [TBD]."),

    # 12. Performance
    item("PERF-01", "12. Performance", "20-player test", "Chạy workload đại diện 20 player với active/unobserved plot, purchase, customer, co-op và visit mix.", "Đảm bảo MVP ổn định ở tải mục tiêu cơ bản.", "PERF-03..PERF-08; TST-08", "P0", "XL", "Kịch bản [TBD] chạy >= 60 phút; MSPT p95 < 50 ms và dưới budget đã chốt; TPS không suy giảm kéo dài; 0 duplicate cash/unlock; entity/task/chunk không tăng vô hạn."),
    item("PERF-02", "12. Performance", "50-player test", "Chạy mixed workload 50 player trên production-equivalent hardware và ghi rõ đây là stretch gate.", "Xác định capacity thực tế và giới hạn an toàn.", "PERF-01; PERF-03..PERF-08", "P0", "XL", "Kịch bản [TBD] chạy >= 60 phút; ghi MSPT median/p95/p99/max, CPU, heap, GC; không crash/deadlock/data corruption; kết quả pass/fail theo budget [TBD] được phê duyệt."),
    item("PERF-03", "12. Performance", "Active plot budget", "Đặt cap dynamic cho active plot, loaded/entity-ticking chunk, display và scheduled work.", "Giữ server ổn định khi nhiều nhà hàng hoạt động.", "CUS-07; CON-01", "P0", "L", "Cap được cấu hình và export metric; vượt cap áp dụng backpressure trong <= [tick]; inactive forced-loaded plot = 0; 1 giờ test không leak chunk ticket/task."),
    item("PERF-04", "12. Performance", "NPC budget", "Giới hạn physical customer theo plot/server; observed ưu tiên, unobserved dùng logical batch.", "Duy trì cảm giác đông khách trong khả năng server.", "CUS-01; CUS-07", "P0", "L", "Normal cap <= [số đã benchmark]/plot và global cap [TBD]; không spawn khi hết budget; overload ceiling chỉ dùng test; entity count trở về baseline sau deactivate."),
    item("PERF-05", "12. Performance", "Pathfinding budget", "Giới hạn request theo tick/plot/server, stagger update và chỉ recalc khi transition/watchdog.", "Giảm spike MSPT do villager navigation.", "CUS-06; PERF-04", "P0", "L", "Path request không vượt cap cấu hình trong stress test; không gọi recalc mỗi tick; queue có backpressure; spark profile pathfinding nằm dưới [budget % CPU/MSPT]."),
    item("PERF-06", "12. Performance", "Entity tracking", "Theo dõi customer, staff presentation, display, Interaction và tracked entities/client.", "Ngăn client/server quá tải vì entity hiển thị.", "CON-09; PERF-04", "P1", "M", "Metric theo type/plot/client cập nhật <= [giây]; visible purchase display < [budget]/active plot; deactivate đưa count về baseline; alert khi vượt threshold [TBD]."),
    item("PERF-07", "12. Performance", "Database latency", "Đo pool wait/query/transaction p50/p95/p99, queue depth và degraded-mode threshold.", "Giữ giao dịch nhanh mà không block main thread.", "PER-01; ECO-09", "P0", "L", "Main-thread SQL/file I/O = 0; p95 transaction < [ms] ở test tải; pool timeout chuyển degraded mode rõ; queue bounded và không mất committed command."),
    item("PERF-08", "12. Performance", "Tick-time monitoring", "Xuất MSPT percentiles, stalls, task cost, active plot và NPC/path count; profile bằng spark.", "Cho operator phát hiện sớm nguyên nhân lag.", "[Monitoring stack TBD]", "P0", "M", "Dashboard có MSPT median/p95/p99/max và correlation với plot/NPC; alert khi vượt [ngưỡng] trong [thời gian]; spark profile thu được cho cả test 20/50 player."),
    item("PERF-09", "12. Performance", "Soak test", "Chạy dài với join/leave, visit, plot churn, purchase, restart nhẹ và database load.", "Phát hiện memory leak và suy giảm chỉ xuất hiện theo thời gian.", "PERF-01; PERF-07; PERF-08", "P0", "XL", "Soak >= [thời lượng, tối thiểu 4h]; heap sau GC, entity, task, chunk ticket và DB connection không có xu hướng tăng ngoài [ngưỡng]; 0 deadlock/data corruption; report được lưu."),

    # 13. Testing
    item("TST-01", "13. Testing", "Unit test", "Bao phủ unlock graph, economy math, state transition, timeout, permission và prestige reset/retain.", "Giảm regression ở logic cốt lõi.", "Các module logic hoàn tất", "P0", "L", "100% invariant P0 có test; line/branch coverage đạt [mục tiêu]; test deterministic chạy 3 lần cùng kết quả; CI thất bại khi có test fail."),
    item("TST-02", "13. Testing", "Integration test", "Test PostgreSQL, Paper lifecycle, world operation, scheduler boundary, inventory và entity cleanup.", "Xác minh các module hoạt động đúng cùng nhau.", "PER-01..PER-08", "P0", "XL", "Suite chạy trên build pin và PostgreSQL thật; tất cả flow critical pass; không synchronous SQL main thread; cleanup sau suite để entity/task/data test dư = 0."),
    item("TST-03", "13. Testing", "Co-op race test", "Chạy concurrent invite, permission, purchase, assignment, dish claim, payout và disconnect.", "Ngăn lỗi chỉ xảy ra khi nhiều người thao tác cùng lúc.", "COP-01..COP-08", "P0", "XL", "Mỗi race scenario chạy >= 100 lần ở 2-4 client; single-winner invariant giữ 100%; reward không vượt budget; không deadlock/stuck assignment."),
    item("TST-04", "13. Testing", "Prestige test", "Kiểm tra gate, confirmation, reset/retain, duplicate request, branch và bonus cap.", "Đảm bảo vòng lặp cuối không phá dữ liệu người chơi.", "PRE-01..PRE-08", "P0", "XL", "Pass 100% prestige matrix; unauthorized request = 0 mutation; duplicate request chỉ tăng 1 level; reset/retain khớp 100%; audit event đầy đủ."),
    item("TST-05", "13. Testing", "Restart test", "Restart ở từng phase purchase, order, cooking, serving, settlement, reset và prestige.", "Bảo vệ tiến trình khi bảo trì hoặc restart bất ngờ.", "PER-06; PRE-07", "P0", "XL", "Mỗi checkpoint chạy >= 10 lần; sau restart không mất/duplicate cash, unlock, dish, payout; world/state hội tụ trong <= [phút]; leak count = 0."),
    item("TST-06", "13. Testing", "Database outage test", "Mô phỏng timeout, disconnect, pool exhaustion và recovery; kiểm tra read-only degraded mode.", "Không cho economy sai khi database gặp sự cố.", "PER-01; PERF-07", "P0", "L", "Khi DB down, purchase/payout/prestige/member/item entitlement mới bị từ chối 100%; không mutate memory-only; player/admin nhận message; recovery không replay trùng command."),
    item("TST-07", "13. Testing", "Anti-exploit test", "Test spam click, replay, inventory move/drop/death, visitor bypass, permission race, AFK macro và stale callback.", "Bảo vệ economy và dữ liệu khỏi exploit phổ biến.", "ECO-09; VIS-03; COP-05; PER-08", "P0", "XL", "100% case trong exploit matrix [TBD] không tạo duplicate cash/unlock/dish/reward; unauthorized mutation = 0; attempt critical có security event; reconciliation chênh lệch = 0."),
    item("TST-08", "13. Testing", "Performance test", "Tự động hóa workload generator, metric capture và pass/fail report cho 20/50 player.", "Cho capacity claim dựa trên bằng chứng lặp lại được.", "PERF-01..PERF-09", "P0", "XL", "Cùng seed/config tái tạo workload sai lệch event rate <= [ngưỡng %]; report có hardware/build/config/MSPT/GC/entity/DB; threshold fail làm pipeline fail; artifact được lưu."),
    item("TST-09", "13. Testing", "Acceptance test", "Chạy end-to-end từ claim đến first prestige ở solo và co-op, gồm visit/admin/recovery.", "Xác nhận MVP Alpha thực sự hoàn chỉnh và có thể chơi.", "Tất cả hạng mục P0; Milestone M5", "P0", "XL", ">= [cỡ mẫu] người mới hoàn thành flow; P50 active time 150-210 phút và co-op floor >= 120 phút; 0 blocker/data corruption; 100% checklist DoD đạt hoặc có waiver được duyệt."),
]


milestones = [
    ["M1", "Foundation & contracts", "PostgreSQL, migration, ledger/idempotency, recovery contract và test harness sẵn sàng.", "[Tuần/mốc TBD]", "PER-01, PER-02, ECO-09, PER-04, TST-01", "Chưa làm", "0/5 deliverable đạt"],
    ["M2", "Progression & content complete", "40-60 node, 2 branch, build stage, content và economy sheet đã cấu hình/validate.", "[Tuần/mốc TBD]", "PRG-01..PRG-07, CON-01..CON-09, ECO-01..ECO-05", "Chưa làm", "Validator 0 lỗi; content lock [TBD]"],
    ["M3", "Multiplayer feature complete", "Customer/staff ổn định, co-op 2-4, visit và prestige hoàn tất.", "[Tuần/mốc TBD]", "CUS-01..CUS-08, STF-01..STF-09, COP-01..COP-08, VIS-01..VIS-06, PRE-01..PRE-08", "Chưa làm", "Critical flow pass; race suite pass"],
    ["M4", "Operational alpha", "Analytics, admin tools, backup/restore, reconciliation và observability hoạt động.", "[Tuần/mốc TBD]", "ANA-01..ANA-08, ADM-01..ADM-08, PER-08, PER-09", "Chưa làm", "Restore drill và admin audit pass"],
    ["M5", "MVP Alpha release gate", "Acceptance, exploit, restart, soak và load test 20/50 player có báo cáo được duyệt.", "[Tuần/mốc TBD]", "PERF-01..PERF-09, TST-01..TST-09", "Chưa làm", "P0 complete; 0 blocker; P50 150-210 phút"],
]

risks = [
    ["R-01", "NPC/pathfinding làm MSPT vượt ngân sách", "Cao", "Cao", "Fixed anchors, global/per-plot budget, stagger update, logical batch cho unobserved.", "[Owner TBD]", "MSPT p95 >= [ngưỡng] hoặc path queue > [ngưỡng]", "Mở"],
    ["R-02", "Economy duplicate do retry/race/restart", "Rất cao", "Trung bình", "Atomic ledger, stable idempotency key, unique constraint, crash/race test.", "[Owner TBD]", "Reconciliation delta != 0", "Mở"],
    ["R-03", "Pacing không đạt 150-210 phút", "Cao", "Cao", "Balance sheet, staged playtest, funnel/bottleneck analytics, tune theo P50/P80.", "[Owner TBD]", "P50 ngoài 150-210 phút", "Mở"],
    ["R-04", "Automation xóa hết active gameplay", "Cao", "Trung bình", "Intervention cadence, 15-25% target, hard cap 30%, đo participation.", "[Owner TBD]", "Active job usage sau automation < [ngưỡng]", "Mở"],
    ["R-05", "Co-op nhân throughput/reward", "Rất cao", "Trung bình", "Fixed order budget, CAS assignment, 1,25x cap, co-op race suite.", "[Owner TBD]", "4-player throughput > 1,25x solo", "Mở"],
    ["R-06", "DB/world lệch sau crash", "Rất cao", "Trung bình", "DB authoritative, canonical projection, fenced operation, reconciliation.", "[Owner TBD]", "Committed stage != world revision", "Mở"],
    ["R-07", "Entity/task/chunk leak trong visit/reset", "Cao", "Trung bình", "Explicit cleanup, ownership tag, lifecycle tests, soak metrics.", "[Owner TBD]", "Count không về baseline sau deactivate", "Mở"],
    ["R-08", "Scope 108 hạng mục vượt khung Phase 2", "Cao", "Cao", "Giữ P0 vertical slice, content placeholder, cắt P2/P3 trước khi giảm hardening.", "[Owner TBD]", "Burn-up lệch > [ngưỡng] trong [số] sprint", "Mở"],
    ["R-09", "50-player gate không đạt trên hardware mục tiêu", "Cao", "Trung bình", "Benchmark sớm, dynamic caps, workload mix, công bố capacity theo số đo.", "[Owner TBD]", "MSPT/GC/DB vượt budget trong PERF-02", "Mở"],
    ["R-10", "Backup restore tạo duplicate hoặc mất progression", "Rất cao", "Thấp", "RPO/RTO, artifact checksum, restore drill, replay stage và reconcile.", "[Owner TBD]", "Restore reconciliation delta != 0", "Mở"],
]

dependencies = [
    ["D-01", "Minecraft/Paper/Java version", "External", "[Chọn exact versions/builds]", "Kiến trúc API, scheduler, resource pack và test matrix", "Trước M1", "[Owner TBD]", "Chưa chốt"],
    ["D-02", "Production-equivalent hardware", "External", "[Chọn CPU/RAM/storage/network]", "Kết quả load test 20/50 player", "Trước PERF-01", "[Owner TBD]", "Chưa chốt"],
    ["D-03", "PostgreSQL hosting/version", "External", "[Chọn provider/version/HA]", "Persistence, latency, backup/restore", "Trước M1", "[Owner TBD]", "Chưa chốt"],
    ["D-04", "Content/stage pipeline", "Internal", "[Bukkit structure hoặc WorldEdit API]", "40-60 node và world projection", "Trước M2", "[Owner TBD]", "Chưa chốt"],
    ["D-05", "Resource pack pipeline", "Internal/External", "[Bắt buộc/tùy chọn; hosting/hash/version TBD]", "Visual, sound, animation và client UX", "Trước content lock", "[Owner TBD]", "Chưa chốt"],
    ["D-06", "Analytics/monitoring stack", "External", "[Chọn sink/dashboard/alert]", "Pacing tune, error detection, performance gate", "Trước M3", "[Owner TBD]", "Chưa chốt"],
    ["D-07", "Permission/admin policy", "Internal", "[Chọn permission model/provider]", "Co-op, visitor, admin tools và audit", "Trước COP-02", "[Owner TBD]", "Chưa chốt"],
    ["D-08", "Backup storage và RPO/RTO", "External/Internal", "[Chọn storage, retention, RPO, RTO]", "Disaster recovery gate", "Trước PER-09", "[Owner TBD]", "Chưa chốt"],
    ["D-09", "Playtest cohort", "External", "[Số người mới, solo/co-op mix TBD]", "Pacing 150-210 phút và acceptance", "Trước M5", "[Owner TBD]", "Chưa chốt"],
    ["D-10", "Content decisions", "Internal", "[Theme, area, recipe, node count, giá, reward TBD]", "Progression/content/economy configuration", "Trước M2", "[Owner TBD]", "Chưa chốt"],
]

dod = [
    ["DoD-01", "Scope", "Toàn bộ P0 hoàn thành; P1 bắt buộc cho Alpha được duyệt; P2/P3 còn lại có quyết định rõ.", "Backlog filter + sign-off", "[Approver TBD]", "Chưa đạt"],
    ["DoD-02", "Gameplay", "Flow claim-to-first-prestige chơi được end-to-end ở solo và co-op 2-4.", "TST-09 pass", "[Approver TBD]", "Chưa đạt"],
    ["DoD-03", "Pacing", "P50 active play đến first prestige nằm 150-210 phút; co-op không dưới 120 phút.", "Playtest dashboard với cỡ mẫu [TBD]", "[Approver TBD]", "Chưa đạt"],
    ["DoD-04", "Progression", "Có 40-60 purchase node hợp lệ và 2 branch không soft-lock.", "Graph validator 0 lỗi + branch playtest", "[Approver TBD]", "Chưa đạt"],
    ["DoD-05", "Data integrity", "Không duplicate/mất cash, unlock, payout, dish hoặc prestige qua race/retry/restart.", "TST-03/04/05/07 pass; reconciliation delta = 0", "[Approver TBD]", "Chưa đạt"],
    ["DoD-06", "Recovery", "Restart/crash recovery và canonical world projection hội tụ không cần sửa DB thủ công.", "PER-06/07/08 acceptance đạt", "[Approver TBD]", "Chưa đạt"],
    ["DoD-07", "Security", "Visitor, collaborator và admin permission được enforce server-side; exploit matrix không có mutation trái phép.", "TST-07 + audit review", "[Approver TBD]", "Chưa đạt"],
    ["DoD-08", "Performance", "Kịch bản 20 player đạt gate; kịch bản 50 player có kết quả pass/fail và capacity decision được duyệt.", "PERF-01/02 report trên hardware mục tiêu", "[Approver TBD]", "Chưa đạt"],
    ["DoD-09", "Stability", "Soak test không leak entity/task/chunk/connection và không có deadlock/data corruption.", "PERF-09 pass", "[Approver TBD]", "Chưa đạt"],
    ["DoD-10", "Operations", "Analytics, structured logs, admin inspect/repair/reconcile và alert critical hoạt động.", "M4 sign-off", "[Approver TBD]", "Chưa đạt"],
    ["DoD-11", "Disaster recovery", "Backup/restore drill đạt RPO/RTO đã chốt và reconciliation chênh lệch bằng 0.", "PER-09 biên bản drill", "[Approver TBD]", "Chưa đạt"],
    ["DoD-12", "Quality gate", "0 blocker/critical bug mở; mọi waiver có owner, expiry và approver; release artifact/version được pin.", "Bug tracker + release checklist", "[Approver TBD]", "Chưa đạt"],
]


def style_table(ws, start_row, headers, data, table_name, widths=None):
    for col, header in enumerate(headers, 1):
        cell = ws.cell(start_row, col, header)
        cell.fill = PatternFill("solid", fgColor="17324D")
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r_idx, row in enumerate(data, start_row + 1):
        for c_idx, value in enumerate(row, 1):
            cell = ws.cell(r_idx, c_idx, value)
            cell.font = Font(name="Arial", size=10, color="1F2933")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=Side(style="hair", color="D9E2EA"))
    end_row = start_row + len(data)
    end_col = len(headers)
    table = Table(displayName=table_name, ref=f"A{start_row}:{get_column_letter(end_col)}{end_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False,
                                          showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    ws.add_table(table)
    if widths:
        for idx, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(idx)].width = width
    return end_row


wb = Workbook()
ws = wb.active
ws.title = "Phase 2 - MVP Alpha"
ws.sheet_view.showGridLines = False

# Title and fill-in legend.
ws.merge_cells("A1:M1")
ws["A1"] = "MINECRAFT RESTAURANT TYCOON | PHASE 2 - MVP ALPHA"
ws["A1"].font = Font(name="Arial", size=18, bold=True, color="FFFFFF")
ws["A1"].fill = PatternFill("solid", fgColor="102A43")
ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
ws.row_dimensions[1].height = 32

ws.merge_cells("A2:M2")
ws["A2"] = "Mục tiêu: mở rộng vertical slice thành MVP Alpha ổn định, chính xác dữ liệu và chống exploit; first prestige P50 trong 150-210 phút active play."
ws["A2"].font = Font(name="Arial", size=10, italic=True, color="334E68")
ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
ws.row_dimensions[2].height = 28

ws.merge_cells("A3:M3")
ws["A3"] = "Cần điền: các placeholder [TBD]/[Chọn ...], Người phụ trách, Trạng thái và Tiến độ %. Ví dụ cập nhật: Owner = Backend Dev; Trạng thái = Đang làm; Tiến độ = 40%."
ws["A3"].fill = PatternFill("solid", fgColor="FFF3BF")
ws["A3"].font = Font(name="Arial", size=10, bold=True, color="5F4B00")
ws["A3"].alignment = Alignment(wrap_text=True, vertical="center")
ws.row_dimensions[3].height = 30

start = 5
end = style_table(ws, start, HEADERS, rows, "Phase2Backlog",
                  [12, 23, 25, 50, 38, 28, 12, 13, 18, 17, 12, 65, 28])
ws.freeze_panes = "A6"
ws.auto_filter.ref = f"A{start}:M{end}"
ws.row_dimensions[start].height = 38
for row_num in range(start + 1, end + 1):
    ws.row_dimensions[row_num].height = 76
    ws.cell(row_num, 11).number_format = '0"%"'
    ws.cell(row_num, 11).alignment = Alignment(horizontal="center", vertical="top")
    for col in (1, 7, 8, 10):
        ws.cell(row_num, col).alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)

# Input styling and validations.
input_fill = PatternFill("solid", fgColor="FFF8DB")
for row_num in range(start + 1, end + 1):
    for col in (9, 10, 11, 13):
        ws.cell(row_num, col).fill = input_fill

lists = wb.create_sheet("_DanhMuc")
for col, values in enumerate((STATUS, PRIORITY, COMPLEXITY), 1):
    for row, value in enumerate(values, 1):
        lists.cell(row, col, value)
lists.sheet_state = "hidden"

status_dv = DataValidation(type="list", formula1="'_DanhMuc'!$A$1:$A$8", allow_blank=False)
priority_dv = DataValidation(type="list", formula1="'_DanhMuc'!$B$1:$B$4", allow_blank=False)
complexity_dv = DataValidation(type="list", formula1="'_DanhMuc'!$C$1:$C$5", allow_blank=False)
progress_dv = DataValidation(type="whole", operator="between", formula1="0", formula2="100", allow_blank=False)
for dv in (status_dv, priority_dv, complexity_dv, progress_dv):
    ws.add_data_validation(dv)
status_dv.add(f"J{start + 1}:J{end}")
priority_dv.add(f"G{start + 1}:G{end}")
complexity_dv.add(f"H{start + 1}:H{end}")
progress_dv.add(f"K{start + 1}:K{end}")

ws.conditional_formatting.add(f"K{start + 1}:K{end}", ColorScaleRule(
    start_type="num", start_value=0, start_color="FEE2E2",
    mid_type="num", mid_value=50, mid_color="FEF3C7",
    end_type="num", end_value=100, end_color="DCFCE7"))
ws.conditional_formatting.add(f"A{start + 1}:M{end}", FormulaRule(
    formula=[f'$J{start + 1}="Bị chặn"'], fill=PatternFill("solid", fgColor="FECACA")))
ws.conditional_formatting.add(f"A{start + 1}:M{end}", FormulaRule(
    formula=[f'$J{start + 1}="Hoàn thành"'], fill=PatternFill("solid", fgColor="DCFCE7")))

# Footer management tables.
section_row = end + 3
sections = [
    ("MILESTONES", ["ID", "Milestone", "Mô tả/Exit scope", "Mốc mục tiêu", "Hạng mục liên quan", "Trạng thái", "Exit metric"], milestones, "MilestonesTable", [12, 28, 65, 20, 48, 16, 36]),
    ("RISKS", ["ID", "Risk", "Impact", "Probability", "Mitigation", "Owner", "Trigger/Metric", "Trạng thái"], risks, "RisksTable", [12, 42, 14, 14, 62, 18, 42, 15]),
    ("DEPENDENCIES", ["ID", "Dependency", "Loại", "Quyết định/Input cần chốt", "Ảnh hưởng", "Cần trước", "Owner", "Trạng thái"], dependencies, "DependenciesTable", [12, 34, 18, 45, 48, 20, 18, 17]),
    ("DEFINITION OF DONE", ["ID", "Nhóm", "Tiêu chí", "Bằng chứng", "Approver", "Trạng thái"], dod, "DefinitionOfDoneTable", [12, 22, 75, 48, 18, 17]),
]

for title, headers, data, table_name, widths in sections:
    ws.merge_cells(start_row=section_row, start_column=1, end_row=section_row, end_column=len(headers))
    title_cell = ws.cell(section_row, 1, title)
    title_cell.fill = PatternFill("solid", fgColor="0B6E69")
    title_cell.font = Font(name="Arial", size=13, bold=True, color="FFFFFF")
    title_cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[section_row].height = 25
    section_row += 1
    table_end = style_table(ws, section_row, headers, data, table_name, widths)
    ws.row_dimensions[section_row].height = 32
    for r in range(section_row + 1, table_end + 1):
        ws.row_dimensions[r].height = 52
    section_row = table_end + 3

ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.print_title_rows = "$1:$5"
ws.sheet_properties.outlinePr.summaryBelow = True

# Workbook-wide professional font.
for sheet in wb.worksheets:
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is not None and cell.font.name != "Arial":
                cell.font = Font(name="Arial", size=cell.font.sz or 10, bold=cell.font.bold,
                                 italic=cell.font.italic, color=cell.font.color)

wb.save(OUTPUT)

# Basic structural verification after serialization.
check = load_workbook(OUTPUT, data_only=False)
assert "Phase 2 - MVP Alpha" in check.sheetnames
sheet = check["Phase 2 - MVP Alpha"]
assert len(rows) == 108
assert sheet["A6"].value == "PRG-01"
assert sheet.cell(end, 1).value == "TST-09"
assert set(STATUS) == {"Chưa làm", "Đang thiết kế", "Sẵn sàng", "Đang làm", "Đang kiểm thử", "Bị chặn", "Hoàn thành", "Hủy"}
assert all(r[11] and len(r[11]) >= 60 for r in rows)
print(f"Created {OUTPUT} with {len(rows)} backlog items and {len(milestones) + len(risks) + len(dependencies) + len(dod)} governance rows.")
